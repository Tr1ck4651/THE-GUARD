from PyQt5.QtWidgets import (
    QApplication,
    QMainWindow,
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QLabel,
    QPushButton,
    QLineEdit,
    QComboBox,
    QTableWidget,
    QTableWidgetItem,
    QHeaderView,
    QDialog,
    QFileDialog,
    QMessageBox,
    QScrollArea,
    QFrame,
    QGroupBox,
    QRadioButton,
    QButtonGroup,
    QGridLayout,
    QFormLayout,
    QDateEdit,
    QTimeEdit,
    QSpinBox,
    QCheckBox,
    QSizePolicy,
    QAbstractItemView,
    QTabWidget,
    QTextEdit,
)
from PyQt5.QtCore import Qt, QDate, QDateTime, QRegExp, QTime, QSize
from PyQt5.QtGui import QIcon, QFont, QPixmap, QColor, QValidator, QRegExpValidator
import os
import json
import uuid 
import re
import datetime
import random
import io
import shutil
import openpyxl

REQUESTS_FILE = "requests.json"

STATUS_TRANSLATIONS = {
    "pending": "В процессе",
    "approved": "Одобрена",
    "rejected": "Отклонена",
}


class ViewEditIndividualRequestDialog(QDialog):
    def __init__(self, request_data, is_editable=True, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Просмотр заявки")
        self.setMinimumSize(900, 700)
        self.request_data = request_data
        self.is_editable = is_editable

        self.setStyleSheet(
            """
            QDialog {
                background: #F4F5F7;
                font-family: 'Segoe UI', 'Roboto', Arial, sans-serif;
            }
            QLabel[section-title] {
                font-size: 22px;
                font-weight: bold;
                color: #FF9800;
                margin-bottom: 18px;
            }
            QFrame[block="true"] {
                background: #fff;
                border-radius: 18px;
                border: 1.5px solid #E0E6ED;
                box-shadow: 0 4px 24px #bbb3;
                margin-bottom: 28px;
            }
            QLabel[key] {
                color: #888;
                font-size: 15px;
                font-weight: 500;
            }
            QLabel[value] {
                color: #222;
                font-size: 16px;
                font-weight: bold;
            }
            QPushButton {
                font-size: 18px;
                padding: 12px 36px;
                border-radius: 12px;
                font-weight: bold;
                background: #FF9800;
                color: white;
                margin: 18px 0 0 0;
            }
            QPushButton:hover {
                background: #e67c00;
            }
        """
        )

        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(30, 30, 30, 30)
        main_layout.setSpacing(0)

        top_panel = QHBoxLayout()
        req_id = request_data.get("id", "-")
        status = request_data.get("status", "pending")
        status_text = STATUS_TRANSLATIONS.get(status, "Неизвестно")
        status_color = (
            "#4CAF50"
            if status == "approved"
            else ("#f44336" if status == "rejected" else "#FF9800")
        )
        title = QLabel(f"Заявка №{req_id}")
        title.setStyleSheet(
            "font-size: 32px; font-weight: bold; color: #FF9800; margin-right: 24px;"
        )
        top_panel.addWidget(title)
        status_badge = QLabel(
            f"<span style='background:{status_color};color:#fff;padding:8px 24px;border-radius:16px;font-size:18px;font-weight:bold;'>{status_text}</span>"
        )
        top_panel.addWidget(status_badge)
        top_panel.addStretch()
        main_layout.addLayout(top_panel)
        main_layout.addSpacing(18)

        info_block = QFrame()
        info_block.setProperty("block", True)
        info_layout = QGridLayout(info_block)
        info_layout.setContentsMargins(32, 24, 32, 24)
        info_layout.setHorizontalSpacing(32)
        info_layout.setVerticalSpacing(10)
        info_block.setMinimumHeight(90)
        info_layout.addWidget(self._key_label("Тип заявки:"), 0, 0)
        info_layout.addWidget(
            self._value_label(
                "Индивидуальная"
                if request_data.get("type") == "individual"
                else "Групповая"
            ),
            0,
            1,
        )
        info_layout.addWidget(self._key_label("Дата создания:"), 0, 2)
        info_layout.addWidget(
            self._value_label(request_data.get("created_at", "-")), 0, 3
        )
        info_layout.addWidget(self._key_label("Действует с:"), 1, 0)
        info_layout.addWidget(
            self._value_label(request_data.get("date_from", "-")), 1, 1
        )
        info_layout.addWidget(self._key_label("Действует по:"), 1, 2)
        info_layout.addWidget(self._value_label(request_data.get("date_to", "-")), 1, 3)
        main_layout.addWidget(info_block)

        host_block = QFrame()
        host_block.setProperty("block", True)
        host_layout = QGridLayout(host_block)
        host_layout.setContentsMargins(32, 24, 32, 24)
        host_layout.setHorizontalSpacing(32)
        host_layout.setVerticalSpacing(10)
        host_block.setMinimumHeight(70)
        host_layout.addWidget(self._key_label("Подразделение:"), 0, 0)
        host_layout.addWidget(
            self._value_label(request_data.get("department", "-")), 0, 1
        )
        host_layout.addWidget(self._key_label("ФИО принимающей стороны:"), 0, 2)
        host_layout.addWidget(
            self._value_label(request_data.get("host_fio", "-")), 0, 3
        )
        main_layout.addWidget(host_block)

        # --- Данные посетителя ---
        visitor_block = QFrame()
        visitor_block.setProperty("block", True)
        visitor_layout = QGridLayout(visitor_block)
        visitor_layout.setContentsMargins(32, 24, 32, 24)
        visitor_layout.setHorizontalSpacing(32)
        visitor_layout.setVerticalSpacing(10)
        visitor_block.setMinimumHeight(120)
        fio = f"{request_data.get('last_name', '')} {request_data.get('first_name', '')} {request_data.get('middle_name', '')}".strip()
        visitor_layout.addWidget(self._key_label("ФИО:"), 0, 0)
        visitor_layout.addWidget(self._value_label(fio), 0, 1)
        visitor_layout.addWidget(self._key_label("Телефон:"), 0, 2)
        visitor_layout.addWidget(
            self._value_label(request_data.get("phone", "-")), 0, 3
        )
        visitor_layout.addWidget(self._key_label("Email:"), 1, 0)
        visitor_layout.addWidget(
            self._value_label(request_data.get("email", "-")), 1, 1
        )
        visitor_layout.addWidget(self._key_label("Организация:"), 1, 2)
        visitor_layout.addWidget(self._value_label(request_data.get("org", "-")), 1, 3)
        visitor_layout.addWidget(self._key_label("Дата рождения:"), 2, 0)
        visitor_layout.addWidget(
            self._value_label(request_data.get("birth", "-")), 2, 1
        )
        visitor_layout.addWidget(self._key_label("Паспорт:"), 2, 2)
        visitor_layout.addWidget(
            self._value_label(
                f"{request_data.get('series', '')} {request_data.get('number', '')}"
            ),
            2,
            3,
        )
        visitor_layout.addWidget(self._key_label("Примечание:"), 3, 0)
        visitor_layout.addWidget(
            self._value_label(request_data.get("note", "-")), 3, 1, 1, 3
        )
        main_layout.addWidget(visitor_block)

        docs_block = QFrame()
        docs_block.setProperty("block", True)
        docs_layout = QHBoxLayout(docs_block)
        docs_layout.setContentsMargins(32, 18, 32, 18)
        docs_layout.setSpacing(18)
        docs_title = QLabel("Документ:")
        docs_title.setProperty("key", True)
        docs_layout.addWidget(docs_title)
        docs_value = QLabel(request_data.get("attached_file_path", "-"))
        docs_value.setProperty("value", True)
        docs_layout.addWidget(docs_value)
        main_layout.addWidget(docs_block)

        btns = QHBoxLayout()
        btns.addStretch()
        close_btn = QPushButton("Закрыть")
        close_btn.setStyleSheet(
            "background:#eee;font-size:18px;padding:12px 36px;border-radius:12px;color:#222;font-weight:bold;"
        )
        close_btn.clicked.connect(self.reject)
        btns.addWidget(close_btn)
        main_layout.addLayout(btns)

    def _key_label(self, text):
        l = QLabel(text)
        l.setProperty("key", True)
        return l

    def _value_label(self, text):
        l = QLabel(text)
        l.setProperty("value", True)
        return l

    def get_status_color(self, status):
        if status == "approved":
            return "green"
        elif status == "rejected":
            return "red"
        else:
            return "orange"

    def attach_file(self):
        fname, _ = QFileDialog.getOpenFileName(
            self,
            "Прикрепить файл",
            "",
            "Документы (*.pdf *.jpg *.png *.jpeg *.doc *.docx)",
        )
        if fname:
            self.attach_btn.setText(os.path.basename(fname))
            self.attach_btn.setProperty("file_path", fname)
        else:
            self.attach_btn.setText("Прикрепить файл")
            self.attach_btn.setProperty("file_path", "")

    def load_photo(self):
        fname, _ = QFileDialog.getOpenFileName(
            self, "Выбрать фото", "", "Изображения (*.png *.jpg *.jpeg)"
        )
        if fname:
            try:
                from PIL import Image
                from PyQt5.QtGui import QPixmap
                import io

                img = Image.open(fname)
                width, height = img.size
                target_ratio = 3 / 4
                img_ratio = width / height

                if abs(img_ratio - target_ratio) > 0.01:
                    if img_ratio > target_ratio:
                        new_width = int(height * target_ratio)
                        left = (width - new_width) // 2
                        img = img.crop((left, 0, left + new_width, height))
                    else:
                        new_height = int(width / target_ratio)
                        top = (height - new_height) // 2
                        img = img.crop((0, top, width, top + new_height))

                img = img.resize(
                    (self.photo_label.width(), self.photo_label.height()), Image.LANCZOS
                )
                buf = io.BytesIO()
                img.save(buf, format="PNG")
                qt_pixmap = QPixmap()
                qt_pixmap.loadFromData(buf.getvalue(), "PNG")
                self.photo_label.setPixmap(qt_pixmap)
                self.photo_label.setProperty("file_path", fname)
                self.photo_label.setText("")
            except Exception as e:
                QMessageBox.warning(self, "Ошибка", f"Ошибка обработки фото: {e}")
                self.photo_label.setText("Фото")
                self.photo_label.setProperty("file_path", "")
                self.photo_label.clear()
        else:
            self.photo_label.setText("Фото")
            self.photo_label.setProperty("file_path", "")
            self.photo_label.clear()

    def get_data(self):
        return {
            "date_from": self.date_from.date().toString("yyyy-MM-dd"),
            "date_to": self.date_to.date().toString("yyyy-MM-dd"),
            "purpose": self.purpose.currentText(),
            "department": self.department.text(),
            "host_fio": self.host_fio.text(),
            "last_name": self.last_name.text(),
            "first_name": self.first_name.text(),
            "middle_name": self.middle_name.text(),
            "phone": self.phone.text(),
            "email": self.email.text(),
            "org": self.org.text(),
            "note": self.note.text(),
            "birth": self.birth.date().toString("yyyy-MM-dd"),
            "series": self.series.text(),
            "number": self.number.text(),
            "photo_path": (
                self.photo_label.property("file_path")
                if self.photo_label.property("file_path")
                else ""
            ),
            "attached_file_path": (
                self.attach_btn.property("file_path")
                if self.attach_btn.property("file_path")
                else ""
            ),
        }

    def view_edit_request_by_id(self, request_id):
        all_requests = self.load_requests()
        req = next((r for r in all_requests if r.get("id", "") == request_id), None)

        if req:
            is_editable = req.get("status", "") == "pending"
            if req.get("type", "") == "individual":
                dlg = ViewEditIndividualRequestDialog(req, is_editable, self)
                result = dlg.exec_()
                if result == 2: 
                    all_requests = [
                        r for r in all_requests if r.get("id", "") != request_id
                    ]
                    self.save_requests(all_requests)
                    self.update_requests_tables()
                    QMessageBox.information(self, "Удалено", "Заявка успешно удалена.")
                    return
                if result == QDialog.Accepted and is_editable:
                   
                    updated_data = dlg.get_data()
             
                    updated_data["id"] = req["id"]
                    updated_data["status"] = req["status"]
                    updated_data["type"] = req["type"]
                    updated_data["created_at"] = req["created_at"]
 
                    req.update(updated_data)
    
                    self.save_requests(all_requests)
          
                    self.update_requests_tables()
                    QMessageBox.information(
                        self, "Сохранено", "Изменения в заявке сохранены."
                    )

            

    def clear_individual_form(self):
        self.pass_id.setText(str(random.randint(100000, 999999)))
        self.temp_pass.setChecked(False)
        self.time_from.setTime(QTime(9, 0))
        self.time_to.setTime(QTime(18, 0))
        self.date_from.setDate(QDate.currentDate())
        self.date_to.setDate(QDate.currentDate())
        self.purpose.setCurrentIndex(0)
        self.purpose_other.clear()
        self.purpose_other.hide()
        self.department.clear()
        self.host_fio.clear()
        self.last_name.clear()
        self.first_name.clear()
        self.middle_name.clear()
        self.phone.clear()
        self.email.clear()
        self.org.clear()
        self.note.clear()
        self.birth.setDate(QDate.currentDate())
        self.series.clear()
        self.number.clear()
        self.photo_label.clear()
        self.photo_label.setText("Фото")
        self.photo_label.setProperty("file_path", "")
        self.attach_btn.setText("Прикрепить файл")
        self.attach_btn.setProperty("file_path", "")

    def submit_individual_form(self):
        errors = []

        if not self.department.text().strip():
            errors.append("Подразделение обязательно")
        if not self.host_fio.text().strip():
            errors.append("ФИО принимающей стороны обязательно")

        if not self.last_name.text().strip():
            errors.append("Фамилия посетителя обязательна")
        if not self.first_name.text().strip():
            errors.append("Имя посетителя обязательно")

        phone = self.phone.text().strip()
        if phone and not re.match(r"\+7 \(\d{3}\) \d{3}-\d{2}-\d{2}", phone):
            errors.append("Телефон должен быть в формате +7 (###) ###-##-##")

        email = self.email.text().strip()
        if not email:
            errors.append("Email посетителя обязателен")
        elif not validate_email(email):
            errors.append("Неверный формат Email")

        if not self.note.toPlainText().strip():
            errors.append("Примечание обязательно")

        birth_date = self.birth.date()
        if not validate_age(birth_date):
            errors.append("Возраст посетителя должен быть не моложе 14 лет")

        series = self.series.text().strip()
        number = self.number.text().strip()
        if not validate_passport(series, number):
            errors.append("Серия паспорта должна содержать 4 цифры, номер - 6 цифр")

        photo_path = self.photo_label.property("file_path")
        if photo_path:
            is_valid, error_msg = validate_photo(photo_path)
            if not is_valid:
                errors.append(error_msg)

        attached_file_path = self.attach_btn.property("file_path")
        if not attached_file_path:
            errors.append("Скан паспорта обязателен")
        elif not attached_file_path.lower().endswith(".jpg"):
            errors.append("Скан паспорта должен быть в формате JPG")

        if errors:
            QMessageBox.warning(self, "Ошибка валидации", "\n".join(errors))
            return

        request_id = str(uuid.uuid4())
        created_at = QDateTime.currentDateTime().toString(Qt.ISODate)

        req = {
            "id": request_id,
            "type": "individual",
            "status": "pending",
            "created_at": created_at,
            "pass_id": self.pass_id.text(),
            "date_from": self.date_from.date().toString("yyyy-MM-dd"),
            "date_to": self.date_to.date().toString("yyyy-MM-dd"),
            "purpose": (
                self.purpose.currentText()
                if self.purpose.currentText() != "Другое"
                else self.purpose_other.text()
            ),
            "department": self.department.text(),
            "host_fio": self.host_fio.text(),
            "last_name": self.last_name.text(),
            "first_name": self.first_name.text(),
            "middle_name": self.middle_name.text(),
            "phone": self.phone.text(),
            "email": self.email.text(),
            "org": self.org.text(),
            "note": self.note.toPlainText(),
            "birth": self.birth.date().toString("yyyy-MM-dd"),
            "series": self.series.text(),
            "number": self.number.text(),
            "photo_path": (
                self.photo_label.property("file_path")
                if self.photo_label.property("file_path")
                else ""
            ),
            "attached_file_path": (
                self.attach_btn.property("file_path")
                if self.attach_btn.property("file_path")
                else ""
            ),
        }

        data = self.parent().load_requests()
        data.append(req)
        self.parent().save_requests(data)

        self.parent().update_requests_tables()
        if hasattr(self.parent(), "history_table"):
            self.parent().apply_history_filters(
                QDate.currentDate().addDays(-30),  
                QDate.currentDate(),
                "Все",
            )

        QMessageBox.information(
            self,
            "Успех",
            f"Заявка №{request_id} успешно отправлена администратору на одобрение!",
        )
        self.clear_individual_form()
        self.accept()


class ViewEditGroupRequestDialog(QDialog):
    def __init__(self, request_data, is_editable=True, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Просмотр/Редактирование групповой заявки")
        self.setMinimumSize(1200, 800)
        self.request_data = request_data
        self.is_editable = is_editable

        main_layout = QVBoxLayout(self)
        self.setLayout(main_layout)

        self.setStyleSheet(
            """
            QDialog {
                background: #F7F8FA;
                font-family: 'Segoe UI', 'Roboto', Arial, sans-serif;
            }
            QFrame[card="true"] {
                background: #fff;
                border-radius: 32px;
                border: 1.5px solid #E0E6ED;
                box-shadow: 0 8px 32px #bbb3;
                margin-bottom: 32px;
            }
            QLabel[role="block-title"] {
                background: #FF9800;
                color: #fff;
                font-size: 22px;
                font-weight: bold;
                border-radius: 16px;
                padding: 12px 36px;
                margin-bottom: 24px;
            }
            QLineEdit, QComboBox, QDateEdit {
                background: #F8F9FA;
                border-radius: 10px;
                border: 2px solid #E0E6ED;
                padding: 14px;
                font-size: 17px;
                color: #222;
                margin-bottom: 12px;
            }
            QLineEdit:focus, QComboBox:focus, QDateEdit:focus {
                border-color: #FF9800;
                background: #fffbe6;
            }
            QPushButton {
                font-size: 18px;
                padding: 14px 36px;
                border-radius: 14px;
                font-weight: bold;
                background: #FF9800;
                color: white;
                margin: 12px 0;
                transition: background 0.3s;
            }
            QPushButton:hover {
                background: #e67c00;
            }
            QPushButton[secondary="true"] {
                background: #eee;
                color: #222;
            }
            QPushButton[secondary="true"]:hover {
                background: #e0e0e0;
            }
            QTableWidget {
                background: white;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 15px;
                gridline-color: #E0E6ED;
            }
            QHeaderView::section {
                background: #FF9800;
                color: #fff;
                font-size: 16px;
                font-weight: bold;
                padding: 12px;
                border: none;
                border-bottom: 2px solid #E0E6ED;
            }
            QTableWidget::item {
                padding: 12px;
                border-bottom: 1px solid #E0E6ED;
            }
            QTableWidget::item:selected {
                background-color: #fff3e0;
                color: #FF9800;
            }
        """
        )

        top_panel = QFrame()
        top_panel.setStyleSheet(
            "background: qlineargradient(x1:0, y1:0, x2:1, y2:1, stop:0 #7BA89F, stop:1 #FF9800);"
            "border-top-left-radius: 22px; border-top-right-radius: 22px;"
            "min-height: 80px; box-shadow: 0 4px 24px #ff980055;"
        )
        top_panel_layout = QHBoxLayout(top_panel)
        top_panel_layout.setContentsMargins(32, 18, 32, 18)
        icon = QLabel()
        icon.setPixmap(QIcon("group.png").pixmap(48, 48))
        icon.setFixedSize(56, 56)
        icon.setStyleSheet("margin-right: 18px;")
        top_panel_layout.addWidget(icon)
        title_lbl = QLabel(
            "<span style='font-size:32px;font-weight:bold;color:#fff;'>Групповая заявка</span>"
        )
        title_lbl.setStyleSheet("margin-right: 32px;")
        top_panel_layout.addWidget(title_lbl)
        top_panel_layout.addStretch()
        status = request_data.get("status", "pending")
        status_text = STATUS_TRANSLATIONS.get(status, "Неизвестно")
        status_color = (
            "orange"
            if status == "pending"
            else ("green" if status == "approved" else "red")
        )
        status_badge = QLabel(
            f"<span style='background:{status_color};color:#fff;padding:8px 24px;border-radius:16px;font-size:20px;font-weight:bold;'>{status_text}</span>"
        )
        top_panel_layout.addWidget(status_badge)
        main_layout.addWidget(top_panel)

        form_card = QFrame()
        form_card.setStyleSheet(
            """
            margin-top: 40px;
            background: #fff;
            border-radius: 32px;
            """
        )
        form_layout = QVBoxLayout(form_card)
        form_layout.setContentsMargins(80, 56, 80, 56)

        block1 = QHBoxLayout()
        pass_frame = QFrame()
        pass_frame.setStyleSheet(
            """
            background: #F8F9FA;
            border-radius: 24px;
            border: 2px solid #E0E6ED;
            """
        )
        pass_layout = QVBoxLayout(pass_frame)
        pass_title = QLabel("Информация для пропуска")
        pass_title.setProperty("role", "block-title")
        pass_title.setStyleSheet(
            """
            background: #FF9800;
            color: #fff;
            font-size: 22px;
            font-weight: bold;
            border-radius: 16px;
            padding: 12px 36px;
            margin-bottom: 24px;
            """
        )
        pass_layout.addWidget(pass_title)
        pass_grid = QGridLayout()
        pass_grid.setHorizontalSpacing(32)
        pass_grid.setVerticalSpacing(18)

        self.pass_id = QLineEdit(request_data.get("pass_id", ""))
        self.pass_id.setReadOnly(True)
        self.pass_id.setStyleSheet(
            "background: #F5F7FA; color: #888; font-weight: bold; border-radius: 10px; font-size: 18px;"
        )
        pass_grid.addWidget(QLabel("ID пропуска:"), 0, 0)
        pass_grid.addWidget(self.pass_id, 0, 1)

        self.temp_pass = QCheckBox("Временный пропуск")
        self.temp_pass.setStyleSheet(
            """
            QCheckBox {
                font-size: 16px;
                padding: 12px;
            }
            QCheckBox::indicator {
                width: 20px;
                height: 20px;
            }
        """
        )
        pass_grid.addWidget(self.temp_pass, 0, 2, 1, 2)
        self.date_from = QDateEdit(
            QDate.fromString(request_data.get("date_from", ""), "yyyy-MM-dd")
        )
        self.date_from.setCalendarPopup(True)
        self.date_from.setReadOnly(not is_editable)
        self.date_to = QDateEdit(
            QDate.fromString(request_data.get("date_to", ""), "yyyy-MM-dd")
        )
        self.date_to.setCalendarPopup(True)
        self.date_to.setReadOnly(not is_editable)
        pass_grid.addWidget(QLabel("Срок действия заявки: с"), 1, 0)
        pass_grid.addWidget(self.date_from, 1, 1)
        pass_grid.addWidget(QLabel("по"), 1, 2)
        pass_grid.addWidget(self.date_to, 1, 3)
        self.purpose = QLineEdit(request_data.get("purpose", ""))
        self.purpose.setReadOnly(True)
        pass_grid.addWidget(QLabel("Цель посещения:"), 2, 0)
        pass_grid.addWidget(self.purpose, 2, 1, 1, 3)
        pass_layout.addLayout(pass_grid)
        block1.addWidget(pass_frame)

        host_frame = QFrame()
        host_frame.setStyleSheet(
            """
            background: #F8F9FA;
            border-radius: 24px;
            border: 2px solid #E0E6ED;
            """
        )
        host_layout = QVBoxLayout(host_frame)
        host_title = QLabel("Принимающая сторона")
        host_title.setProperty("role", "block-title")
        host_title.setStyleSheet(
            """
            background: #FF9800;
            color: #fff;
            font-size: 22px;
            font-weight: bold;
            border-radius: 16px;
            padding: 12px 36px;
            margin-bottom: 24px;
            """
        )
        host_layout.addWidget(host_title)
        self.department = QLineEdit(request_data.get("department", ""))
        self.department.setReadOnly(True)
        self.host_fio = QLineEdit(request_data.get("host_fio", ""))
        self.host_fio.setReadOnly(True)
        host_layout.addWidget(self.department)
        host_layout.addWidget(self.host_fio)
        block1.addWidget(host_frame)
        form_layout.addLayout(block1)

        guests_block = QFrame()
        guests_block.setProperty("block", True)
        guests_layout = QVBoxLayout(guests_block)
        guests_title = QLabel("Список гостей")
        guests_title.setStyleSheet(
            "font-size: 20px; font-weight: bold; color: #FF9800; margin-bottom: 12px;"
        )
        guests_layout.addWidget(guests_title)
        guests_table = QTableWidget()
        guests_table.setColumnCount(5)
        guests_table.setHorizontalHeaderLabels(
            ["ФИО", "Контакты", "Дата рождения", "Паспорт", "Фото"]
        )
        guests_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        visitors = self.request_data.get("visitors", [])
        guests_table.setRowCount(len(visitors))
        for i, visitor in enumerate(visitors):
            fio = visitor.get("fio", "")
            contacts = visitor.get("contacts", "")
            birth = visitor.get("birth", "")
            passport = visitor.get("passport", "")
            photo = visitor.get("photo_path", "")
            guests_table.setItem(i, 0, QTableWidgetItem(fio))
            guests_table.setItem(i, 1, QTableWidgetItem(contacts))
            guests_table.setItem(i, 2, QTableWidgetItem(birth))
            guests_table.setItem(i, 3, QTableWidgetItem(passport))
            guests_table.setItem(i, 4, QTableWidgetItem(photo))
        guests_layout.addWidget(guests_table)
        form_layout.addWidget(guests_block)

        docs_card = QFrame()
        docs_card.setStyleSheet(
            """
            background: #F8F9FA;
            border-radius: 24px;
            border: 2px solid #E0E6ED;
            """
        )
        docs_layout = QHBoxLayout(docs_card)
        docs_icon = QLabel()
        docs_icon.setPixmap(QIcon("attach.png").pixmap(22, 22))
        docs_icon.setFixedSize(28, 28)
        docs_layout.addWidget(docs_icon)
        docs_fields = QVBoxLayout()
        docs_title = QLabel("<b style='font-size:18px;color:#FF9800;'>Документы</b>")
        docs_fields.addWidget(docs_title)
        self.group_attach_btn = QPushButton()
        self.group_attach_btn.setIcon(QIcon("attach.png"))
        self.group_attach_btn.setText("Прикрепить файл")
        self.group_attach_btn.setStyleSheet(
            "font-size:15px;padding:8px 12px;border-radius:8px;background:#FF9800;color:white;font-weight:bold;"
        )
        self.group_attach_btn.clicked.connect(self.attach_group_file)
        docs_fields.addWidget(self.group_attach_btn)
        docs_layout.addLayout(docs_fields)
        form_layout.addWidget(docs_card)

        btns = QHBoxLayout()
        clear_btn = QPushButton("Очистить форму")
        clear_btn.setStyleSheet(
            "font-size:20px;padding:16px 36px;border-radius:12px;background:#eee;color:#222;font-weight:bold;"
        )
        clear_btn.clicked.connect(self.clear_group_form)
        submit_btn = QPushButton("Оформить заявку")
        submit_btn.setStyleSheet(
            "font-size:20px;padding:16px 36px;border-radius:12px;font-weight:bold;background:#FF9800;color:white;transition: background 0.3s;"
        )
        submit_btn.clicked.connect(self.submit_group_form)
        submit_btn.setCursor(Qt.PointingHandCursor)
        clear_btn.setCursor(Qt.PointingHandCursor)
        btns.addWidget(clear_btn)
        btns.addWidget(submit_btn)
        form_layout.addLayout(btns)

        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)
        scroll_area.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        scroll_area.setWidget(form_card)
        scroll_area.setStyleSheet("border: none; background: transparent;")
        main_layout.addWidget(scroll_area)

    def view_edit_request_by_id(self, request_id):
        all_requests = self.load_requests()
        req = next((r for r in all_requests if r.get("id", "") == request_id), None)

        if req:
            is_editable = req.get("status", "") == "pending"
            if req.get("type", "") == "individual":
                dlg = ViewEditIndividualRequestDialog(req, is_editable, self)
                result = dlg.exec_()
                if result == 2: 
                    all_requests = [
                        r for r in all_requests if r.get("id", "") != request_id
                    ]
                    self.save_requests(all_requests)
                    self.update_requests_tables()
                    QMessageBox.information(self, "Удалено", "Заявка успешно удалена.")
                    return
                if result == QDialog.Accepted and is_editable:
 
                    updated_data = dlg.get_data()
     
                    updated_data["id"] = req["id"]
                    updated_data["status"] = req["status"]
                    updated_data["type"] = req["type"]
                    updated_data["created_at"] = req["created_at"]
        
                    req.update(updated_data)
        
                    self.save_requests(all_requests)
           
                    self.update_requests_tables()
                    QMessageBox.information(
                        self, "Сохранено", "Изменения в заявке сохранены."
                    )

            elif req.get("type", "") == "group":
                dlg = ViewEditGroupRequestDialog(req, is_editable, self)
                dlg.exec_()

    def clear_group_form(self):
        self.group_pass_id.setText(str(random.randint(100000, 999999)))
        self.group_date_from.setDate(QDate.currentDate())
        self.group_date_to.setDate(QDate.currentDate())
        self.group_purpose.setCurrentIndex(0)
        self.group_purpose_other.clear()
        self.group_purpose_other.hide()
        self.department.clear()
        self.host_fio.clear()
        self.visitors_table.setRowCount(0)

    def submit_group_form(self):
        errors = []

        if not self.department.text().strip():
            errors.append("Подразделение обязательно")
        if not self.host_fio.text().strip():
            errors.append("ФИО принимающей стороны обязательно")

        if self.visitors_table.rowCount() == 0:
            errors.append("Добавьте хотя бы одного посетителя в группу")

        if errors:
            QMessageBox.warning(self, "Ошибка", "\n".join(errors))
            return

        try:
            request_id = str(uuid.uuid4())
            pass_id = str(random.randint(100000, 999999))
            created_at = QDateTime.currentDateTime().toString(Qt.ISODate)

            visitors = []
            for row in range(self.visitors_table.rowCount()):
                visitor = {
                    "fio": (
                        self.visitors_table.item(row, 0).text()
                        if self.visitors_table.item(row, 0)
                        else ""
                    ),
                    "contacts": (
                        self.visitors_table.item(row, 1).text()
                        if self.visitors_table.item(row, 1)
                        else ""
                    ),
                    "birth": (
                        self.visitors_table.item(row, 2).text()
                        if self.visitors_table.item(row, 2)
                        else ""
                    ),
                    "passport": (
                        self.visitors_table.item(row, 3).text()
                        if self.visitors_table.item(row, 3)
                        else ""
                    ),
                    "photo_path": (
                        self.visitors_table.item(row, 4).text()
                        if self.visitors_table.item(row, 4)
                        else ""
                    ),
                }
                visitors.append(visitor)

            req = {
                "id": request_id,
                "type": "group",
                "status": "pending",
                "created_at": created_at,
                "pass_id": pass_id,
                "date_from": self.group_date_from.date().toString("yyyy-MM-dd"),
                "date_to": self.group_date_to.date().toString("yyyy-MM-dd"),
                "purpose": (
                    self.group_purpose_other.text()
                    if self.group_purpose.currentText() == "Другое"
                    else self.group_purpose.currentText()
                ),
                "department": self.department.text(),
                "host_fio": self.host_fio.text(),
                "visitors": visitors,
                "attached_file_path": (
                    self.attach_btn.property("file_path")
                    if self.attach_btn.property("file_path")
                    else ""
                ),
            }

            data = self.parent().load_requests()
            if not isinstance(data, list):
                data = []

            data.append(req)

            self.parent().save_requests(data)

            if hasattr(self.parent(), "show_requests"):
                self.parent().show_requests()

            QMessageBox.information(
                self,
                "Успех",
                f"Групповая заявка №{pass_id} успешно отправлена администратору на одобрение!",
            )
            self.clear_group_form()
            self.accept()

        except Exception as e:
            QMessageBox.critical(
                self, "Ошибка", f"Произошла ошибка при сохранении заявки: {str(e)}"
            )
            print(f"Ошибка при сохранении групповой заявки: {e}")

    def show_requests(self):
        self.clear_work_area()

        title = QLabel("Заявки на посещение")
        title.setStyleSheet(
            """
            QLabel {
                font-size: 32px;
                font-weight: bold;
                color: #2c3e50;
                margin-bottom: 20px;
            }
        """
        )
        self.work_layout.addWidget(title)

        self.request_tabs = QTabWidget()
        self.request_tabs.setStyleSheet(
            """
            QTabWidget::pane {
                border: 0;
                background: #F3F4F6;
                border-radius: 16px;
                margin-top: 12px;
            }
            QTabBar::tab {
                background: #fff;
                color: #000;
                padding: 12px 24px;
                border-top-left-radius: 12px;
                border-top-right-radius: 12px;
                margin-right: 4px;
                font-size: 16px;
                font-weight: bold;
                min-width: 160px;
            }
            QTabBar::tab:selected {
                background: #FF9800;
                color: #000;
            }
            QTabBar::tab:hover:!selected {
                background: #ffe0b2;
                color: #000;
            }
        """
        )

        all_requests_tab = QWidget()
        all_requests_layout = QVBoxLayout(all_requests_tab)
        self.all_requests_table = self.create_requests_table()
        all_requests_layout.addWidget(self.all_requests_table)
        self.request_tabs.addTab(all_requests_tab, "Все заявки")

        approved_requests_tab = QWidget()
        approved_requests_layout = QVBoxLayout(approved_requests_tab)
        self.approved_requests_table = self.create_requests_table()
        approved_requests_layout.addWidget(self.approved_requests_table)
        self.request_tabs.addTab(approved_requests_tab, "Одобренные")

        rejected_requests_tab = QWidget()
        rejected_requests_layout = QVBoxLayout(rejected_requests_tab)
        self.rejected_requests_table = self.create_requests_table()
        rejected_requests_layout.addWidget(self.rejected_requests_table)
        self.request_tabs.addTab(rejected_requests_tab, "Отклоненные")

        pending_requests_tab = QWidget()
        pending_requests_layout = QVBoxLayout(pending_requests_tab)
        self.pending_requests_table = self.create_requests_table()
        pending_requests_layout.addWidget(self.pending_requests_table)
        self.request_tabs.addTab(pending_requests_tab, "В процессе")

        history_tab = QWidget()
        history_layout = QVBoxLayout(history_tab)

        filter_layout = QHBoxLayout()

        date_filter_label = QLabel("Период:")
        date_filter_label.setStyleSheet("font-size: 14px; color: #666;")
        date_from = QDateEdit()
        date_from.setCalendarPopup(True)
        date_from.setDisplayFormat("dd.MM.yyyy")
        date_from.setDate(QDate.currentDate().addMonths(-1))
        date_to = QDateEdit()
        date_to.setCalendarPopup(True)
        date_to.setDisplayFormat("dd.MM.yyyy")
        date_to.setDate(QDate.currentDate())

        type_filter_label = QLabel("Тип заявки:")
        type_filter_label.setStyleSheet("font-size: 14px; color: #666;")
        type_filter = QComboBox()
        type_filter.addItems(["Все", "Индивидуальная", "Групповая"])

        apply_filter_btn = QPushButton("Применить")
        apply_filter_btn.setStyleSheet(
            """
            QPushButton {
                background-color: #FF9800;
                color: white;
                border: none;
                padding: 8px 16px;
                border-radius: 8px;
                font-weight: bold;
            }
            QPushButton:hover {
                background-color: #F57C00;
            }
        """
        )

        filter_layout.addWidget(date_filter_label)
        filter_layout.addWidget(date_from)
        filter_layout.addWidget(QLabel("—"))
        filter_layout.addWidget(date_to)
        filter_layout.addSpacing(20)
        filter_layout.addWidget(type_filter_label)
        filter_layout.addWidget(type_filter)
        filter_layout.addSpacing(20)
        filter_layout.addWidget(apply_filter_btn)
        filter_layout.addStretch()

        history_layout.addLayout(filter_layout)

        self.history_table = self.create_history_table()
        history_layout.addWidget(self.history_table)

        self.request_tabs.addTab(history_tab, "История")

        self.work_layout.addWidget(self.request_tabs)

        self.update_requests_tables()

        apply_filter_btn.clicked.connect(
            lambda: self.apply_history_filters(
                date_from.date(), date_to.date(), type_filter.currentText()
            )
        )

    def create_history_table(self):

        table = QTableWidget()
        table.setColumnCount(8)
        table.setHorizontalHeaderLabels(
            [
                "№",
                "Тип заявки",
                "ФИО посетителя",
                "Дата создания",
                "Действует с",
                "Действует по",
                "Статус",
                "Действия",
            ]
        )

        table.setMinimumHeight(600)
   
        table.verticalHeader().setDefaultSectionSize(50)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setSelectionMode(QAbstractItemView.SingleSelection)
        table.setStyleSheet(
            """
            QTableWidget {
                background-color: white;
                border: 1px solid #e0e0e0;
                border-radius: 12px;
                gridline-color: #f0f0f0;
            }
            QTableWidget::item {
                padding: 12px;
                border-bottom: 1px solid #f0f0f0;
            }
            QHeaderView::section {
                background-color: #f8f9fa;
                padding: 16px;
                border: none;
                border-bottom: 2px solid #e0e0e0;
                font-weight: bold;
                color: #2c3e50;
            }
            QTableWidget::item:selected {
                background-color: #fff3e0;
                color: #FF9800;
            }
            QPushButton {
                background-color: #fff;
                color: #FF9800;
                border: 2px solid #FF9800;
                padding: 8px 24px;
                border-radius: 8px;
                font-size: 16px;
                min-width: 120px;
                font-weight: bold;
            }
            QPushButton:hover {
                background-color: #FF9800;
                color: #fff;
            }
        """
        )
        table.setColumnWidth(7, 180)  
        return table

    def apply_history_filters(self, date_from, date_to, request_type):

        requests = self.load_requests()
        filtered_requests = []

        for request in requests:

            created_at = request.get("created_at", "")
            if not created_at:
                continue

            try:
                created_date = QDateTime.fromString(created_at, Qt.ISODate).date()
                if not created_date.isValid():
                    continue

                if date_from <= created_date <= date_to:

                    if (
                        request_type == "Все"
                        or (
                            request_type == "Индивидуальная"
                            and request.get("type") == "individual"
                        )
                        or (
                            request_type == "Групповая"
                            and request.get("type") == "group"
                        )
                    ):
                        filtered_requests.append(request)
            except Exception as e:
                print(f"Ошибка при обработке даты: {e}")
                continue

        self.populate_history_table(filtered_requests)

    def populate_history_table(self, requests_list):

        self.history_table.setRowCount(len(requests_list))
        from PyQt5.QtGui import QIcon

        for row, request in enumerate(requests_list):
  
            id_item = QTableWidgetItem(str(request.get("id", "")))
            self.history_table.setItem(row, 0, id_item)

            type_text = (
                "Индивидуальная" if request.get("type") == "individual" else "Групповая"
            )
            type_item = QTableWidgetItem(type_text)
            self.history_table.setItem(row, 1, type_item)

            fio = f"{request.get('last_name', '')} {request.get('first_name', '')} {request.get('middle_name', '')}".strip()
            fio_item = QTableWidgetItem(fio)
            self.history_table.setItem(row, 2, fio_item)

            created_item = QTableWidgetItem(request.get("created_at", ""))
            self.history_table.setItem(row, 3, created_item)

            date_from_item = QTableWidgetItem(request.get("date_from", ""))
            self.history_table.setItem(row, 4, date_from_item)

            date_to_item = QTableWidgetItem(request.get("date_to", ""))
            self.history_table.setItem(row, 5, date_to_item)

            status = request.get("status", "pending")
            status_text = {
                "approved": "Одобрена",
                "rejected": "Отклонена",
                "pending": "В процессе",
            }.get(status, "В процессе")

            status_item = QTableWidgetItem(status_text)
            status_item.setForeground(self.get_status_color_brush(status))
            self.history_table.setItem(row, 6, status_item)

            view_btn = QPushButton("Просмотреть")
            view_btn.setStyleSheet(
                """
                QPushButton {
                    font-size: 14px;
                    font-weight: bold;
                    color: #000;
                    background: transparent;
                    border: none;
                    padding: 4px 8px;
                }
                QPushButton:hover {
                    color: #FF9800;
                    text-decoration: underline;
                }
                """
            )
            view_btn.setProperty("action-btn", True)
            view_btn.setCursor(Qt.PointingHandCursor)
            view_btn.setIcon(QIcon.fromTheme("search") or QIcon("icons/search.png"))
            view_btn.setIconSize(QSize(20, 20))
            view_btn.clicked.connect(
                lambda checked, r=request: self.view_request_details(r)
            )

            cell_widget = QWidget()
            layout = QHBoxLayout(cell_widget)
            layout.addWidget(view_btn)
            layout.setAlignment(Qt.AlignCenter)
            layout.setContentsMargins(0, 0, 0, 0)
            cell_widget.setLayout(layout)
            cell_widget.setStyleSheet("background: #FFF3E0; border-radius: 12px;")
            self.history_table.setCellWidget(row, 7, cell_widget)

    def clear_work_area(self):
        while self.work_layout.count():
            item = self.work_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()

    def show_create_request_tiles(self):
        self.clear_work_area()

        main_container = QWidget()
        main_layout = QVBoxLayout(main_container)
        main_layout.setAlignment(Qt.AlignCenter)
        main_layout.setContentsMargins(50, 50, 50, 50)
        main_layout.setSpacing(40)

        title = QLabel("Создание заявки")
        title.setStyleSheet(
            """
            QLabel {
                font-size: 36px;
                font-weight: bold;
                color: #2c3e50;
                margin-bottom: 40px;
                text-align: center;
            }
        """
        )
        title.setAlignment(Qt.AlignCenter)
        main_layout.addWidget(title)

        tiles_container = QWidget()
        tiles_layout = QHBoxLayout(tiles_container)
        tiles_layout.setSpacing(40)
        tiles_layout.setContentsMargins(0, 0, 0, 0)
        tiles_layout.setAlignment(Qt.AlignCenter)

        individual_tile = QFrame()
        individual_tile.setObjectName("individualTile")
        individual_tile.setStyleSheet(
            """
            QFrame#individualTile {
                background-color: white;
                border-radius: 24px;
                border: 2px solid #e0e0e0;
                min-width: 450px;
                min-height: 550px;
                transition: all 0.3s ease;
            }
            QFrame#individualTile:hover {
                border-color: #FF9800;
                box-shadow: 0 8px 16px rgba(255, 152, 0, 0.15);
                transform: translateY(-5px);
            }
            QLabel {
                color: #2c3e50;
            }
            QPushButton {
                background-color: #FF9800;
                color: white;
                border: none;
                padding: 16px 32px;
                border-radius: 12px;
                font-size: 18px;
                font-weight: bold;
                min-width: 240px;
                transition: all 0.3s ease;
            }
            QPushButton:hover {
                background-color: #F57C00;
                transform: scale(1.05);
            }
        """
        )
        individual_layout = QVBoxLayout(individual_tile)
        individual_layout.setContentsMargins(40, 40, 40, 40)
        individual_layout.setSpacing(30)
        individual_layout.setAlignment(Qt.AlignCenter)

        individual_icon = QLabel()
        individual_icon.setPixmap(
            QPixmap("icons/individual.png").scaled(
                140, 140, Qt.KeepAspectRatio, Qt.SmoothTransformation
            )
        )
        individual_icon.setAlignment(Qt.AlignCenter)
        individual_layout.addWidget(individual_icon)

        individual_title = QLabel("Индивидуальная заявка")
        individual_title.setStyleSheet(
            "font-size: 28px; font-weight: bold; color: #2c3e50;"
        )
        individual_title.setAlignment(Qt.AlignCenter)
        individual_layout.addWidget(individual_title)

        individual_desc = QLabel(
            "Создание заявки для одного посетителя.\n"
            "Включает персональные данные и документы."
        )
        individual_desc.setStyleSheet("font-size: 18px; color: #666; line-height: 1.5;")
        individual_desc.setAlignment(Qt.AlignCenter)
        individual_desc.setWordWrap(True)
        individual_layout.addWidget(individual_desc)

        individual_layout.addStretch()

        individual_btn = QPushButton("Создать заявку")
        individual_btn.setCursor(Qt.PointingHandCursor)
        individual_btn.clicked.connect(self.open_individual_request_window)
        individual_layout.addWidget(individual_btn, alignment=Qt.AlignCenter)

        group_tile = QFrame()
        group_tile.setObjectName("groupTile")
        group_tile.setStyleSheet(
            """
            QFrame#groupTile {
                background-color: white;
                border-radius: 24px;
                border: 2px solid #e0e0e0;
                min-width: 450px;
                min-height: 550px;
                transition: all 0.3s ease;
            }
            QFrame#groupTile:hover {
                border-color: #FF9800;
                box-shadow: 0 8px 16px rgba(255, 152, 0, 0.15);
                transform: translateY(-5px);
            }
            QLabel {
                color: #2c3e50;
            }
            QPushButton {
                background-color: #FF9800;
                color: white;
                border: none;
                padding: 16px 32px;
                border-radius: 12px;
                font-size: 18px;
                font-weight: bold;
                min-width: 240px;
                transition: all 0.3s ease;
            }
            QPushButton:hover {
                background-color: #F57C00;
                transform: scale(1.05);
            }
        """
        )
        group_layout = QVBoxLayout(group_tile)
        group_layout.setContentsMargins(40, 40, 40, 40)
        group_layout.setSpacing(30)
        group_layout.setAlignment(Qt.AlignCenter)

        group_icon = QLabel()
        group_icon.setPixmap(
            QPixmap("icons/group.png").scaled(
                140, 140, Qt.KeepAspectRatio, Qt.SmoothTransformation
            )
        )
        group_icon.setAlignment(Qt.AlignCenter)
        group_layout.addWidget(group_icon)

        group_title = QLabel("Групповая заявка")
        group_title.setStyleSheet("font-size: 28px; font-weight: bold; color: #2c3e50;")
        group_title.setAlignment(Qt.AlignCenter)
        group_layout.addWidget(group_title)

        group_desc = QLabel(
            "Создание заявки для группы посетителей.\n"
            "Возможность загрузки списка из файла."
        )
        group_desc.setStyleSheet("font-size: 18px; color: #666; line-height: 1.5;")
        group_desc.setAlignment(Qt.AlignCenter)
        group_desc.setWordWrap(True)
        group_layout.addWidget(group_desc)

        group_layout.addStretch()

        group_btn = QPushButton("Создать заявку")
        group_btn.setCursor(Qt.PointingHandCursor)
        group_btn.clicked.connect(self.open_group_request_window)
        group_layout.addWidget(group_btn, alignment=Qt.AlignCenter)

        tiles_layout.addWidget(individual_tile)
        tiles_layout.addWidget(group_tile)

        main_layout.addWidget(tiles_container)
        main_layout.addStretch()

        self.work_layout.addWidget(main_container)
        self.work_layout.addStretch()

    def open_individual_request_window(self):
        dialog = CreateIndividualRequestWindow(self)
        dialog.exec_()

    def open_group_request_window(self):
        dialog = CreateGroupRequestWindow(self)
        dialog.exec_()

    def attach_group_file(self):
        file_name, _ = QFileDialog.getOpenFileName(
            self, "Выберите файл", "", "All Files (*)"
        )
        if file_name:
            QMessageBox.information(
                self, "Успех", f"Файл {file_name} успешно прикреплен"
            )

    def clear_group_form(self):
        self.group_pass_id.setText(str(random.randint(100000, 999999)))
        self.group_date_from.setDate(QDate.currentDate())
        self.group_date_to.setDate(QDate.currentDate())
        self.group_purpose.setCurrentIndex(0)
        self.group_purpose_other.clear()
        self.group_purpose_other.hide()
        self.department.clear()
        self.host_fio.clear()
        self.visitors_table.setRowCount(0)

    def submit_group_form(self):

        if not self.department.text().strip() or not self.host_fio.text().strip():
            QMessageBox.warning(
                self, "Ошибка", "Пожалуйста, заполните все обязательные поля"
            )
            return

        if self.visitors_table.rowCount() == 0:
            QMessageBox.warning(
                self, "Ошибка", "Пожалуйста, добавьте хотя бы одного гостя"
            )
            return

        QMessageBox.information(self, "Успех", "Групповая заявка успешно отправлена")
        self.accept()

    def add_group_visitor_row_custom(self, fields):

        row_count = self.visitors_table.rowCount()
        if row_count < 8:
            self.visitors_table.insertRow(row_count)
            self.visitors_table.setItem(
                row_count, 0, QTableWidgetItem(fields.get("last_name", ""))
            )
            self.visitors_table.setItem(
                row_count, 1, QTableWidgetItem(fields.get("first_name", ""))
            )
            self.visitors_table.setItem(
                row_count, 2, QTableWidgetItem(fields.get("middle_name", ""))
            )
            self.visitors_table.setItem(
                row_count, 3, QTableWidgetItem(fields.get("phone", ""))
            )
            self.visitors_table.setItem(
                row_count, 4, QTableWidgetItem(fields.get("email", ""))
            )
            self.visitors_table.setItem(
                row_count, 5, QTableWidgetItem(fields.get("org", ""))
            )
            self.visitors_table.setItem(
                row_count, 6, QTableWidgetItem(fields.get("note", ""))
            )
            self.visitors_table.setItem(
                row_count, 7, QTableWidgetItem(fields.get("birth", ""))
            )
            self.visitors_table.setItem(
                row_count, 8, QTableWidgetItem(fields.get("series", ""))
            )
            self.visitors_table.setItem(
                row_count, 9, QTableWidgetItem(fields.get("number", ""))
            )

            delete_btn = QPushButton("Удалить")
            delete_btn.clicked.connect(lambda: self.delete_visitor_row(row_count))
            self.visitors_table.setCellWidget(row_count, 10, delete_btn)
        else:
            QMessageBox.warning(
                self, "Внимание", "Максимальное количество посетителей в группе - 8."
            )

    def delete_visitor_row(self, row):

        self.visitors_table.removeRow(row)

    def save_requests(self, data):
        try:
            with open(REQUESTS_FILE, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print(f"Ошибка при сохранении заявок: {e}")

    def view_request_details(self, request):

        if request.get("type") == "individual":
            dialog = ViewEditIndividualRequestDialog(
                request, is_editable=False, parent=self
            )
        else:
            dialog = ViewEditGroupRequestDialog(request, is_editable=False, parent=self)
        dialog.exec_()

    def open_add_guest_dialog(self):
        dialog = QDialog(self)
        dialog.setWindowTitle("Добавить гостя")
        dialog.setMinimumWidth(400)

        layout = QVBoxLayout(dialog)

        form_layout = QFormLayout()

        last_name = QLineEdit()
        first_name = QLineEdit()
        middle_name = QLineEdit()
        phone = QLineEdit()
        phone.setValidator(PhoneValidator())
        email = QLineEdit()
        org = QLineEdit()
        birth = QDateEdit()
        birth.setCalendarPopup(True)
        birth.setDate(QDate.currentDate().addYears(-18))
        series = QLineEdit()
        series.setValidator(PassportValidator())
        number = QLineEdit()
        number.setValidator(PassportValidator())

        form_layout.addRow("Фамилия:", last_name)
        form_layout.addRow("Имя:", first_name)
        form_layout.addRow("Отчество:", middle_name)
        form_layout.addRow("Телефон:", phone)
        form_layout.addRow("Email:", email)
        form_layout.addRow("Организация:", org)
        form_layout.addRow("Дата рождения:", birth)
        form_layout.addRow("Серия паспорта:", series)
        form_layout.addRow("Номер паспорта:", number)

        layout.addLayout(form_layout)

        buttons = QHBoxLayout()
        save_btn = QPushButton("Сохранить")
        cancel_btn = QPushButton("Отмена")

        buttons.addWidget(save_btn)
        buttons.addWidget(cancel_btn)
        layout.addLayout(buttons)

        save_btn.clicked.connect(dialog.accept)
        cancel_btn.clicked.connect(dialog.reject)

        if dialog.exec_() == QDialog.Accepted:
 
            guest_data = {
                "last_name": last_name.text(),
                "first_name": first_name.text(),
                "middle_name": middle_name.text(),
                "phone": phone.text(),
                "email": email.text(),
                "org": org.text(),
                "birth": birth.date().toString("yyyy-MM-dd"),
                "series": series.text(),
                "number": number.text(),
            }

            self.add_group_visitor_row_custom(guest_data)

    def download_guest_template(self):
        try:
            temp_file = "guest_template.xlsx"

            wb = openpyxl.Workbook()
            ws = wb.active

            headers = [
                "Фамилия",
                "Имя",
                "Отчество",
                "Телефон",
                "Email",
                "Организация",
                "Дата рождения",
                "Серия паспорта",
                "Номер паспорта",
            ]
            ws.append(headers)

            wb.save(temp_file)

            file_path, _ = QFileDialog.getSaveFileName(
                self, "Сохранить шаблон", "guest_template.xlsx", "Excel Files (*.xlsx)"
            )

            if file_path:
                shutil.copy2(temp_file, file_path)
                QMessageBox.information(self, "Успех", "Шаблон успешно сохранен")

            os.remove(temp_file)

        except Exception as e:
            QMessageBox.warning(self, "Ошибка", f"Не удалось создать шаблон: {str(e)}")

    def upload_guest_list(self):
        try:
            from openpyxl import load_workbook
            from PyQt5.QtWidgets import QFileDialog
            from PyQt5.QtCore import QDate
            import re

            file_path, _ = QFileDialog.getOpenFileName(
                self, "Загрузить список гостей", "", "Excel Files (*.xlsx)"
            )
            if not file_path:
                return
            wb = load_workbook(file_path)
            ws = wb.active
            headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]
            norm_map = {i: self._normalize_header(h) for i, h in enumerate(headers)}
            required = [
                "last_name",
                "first_name",
                "middle_name",
                "phone",
                "email",
                "org",
                "note",
                "birth",
                "series",
                "number",
            ]
            if not all(v in norm_map.values() for v in required):
                missing = [v for v in required if v not in norm_map.values()]
                QMessageBox.warning(
                    self, "Ошибка", f"В файле отсутствуют столбцы: {', '.join(missing)}"
                )
                return
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            for row in rows:
                if not row:
                    continue
                fields = {}
                for idx, norm in norm_map.items():
                    if norm:
                        fields[norm] = (
                            str(row[idx]).strip() if row[idx] is not None else ""
                        )
                phone_digits = "".join(filter(str.isdigit, fields.get("phone", "")))
                phone_formatted = ""
                if len(phone_digits) == 10:
                    phone_formatted = f"+7 ({phone_digits[0:3]}) {phone_digits[3:6]}-{phone_digits[6:8]}-{phone_digits[8:10]}"
                fields["phone"] = phone_formatted
                try:
                    d, m, y = map(int, str(fields.get("birth", "")).split("."))
                    birth_date = QDate(y, m, d)
                    fields["birth"] = birth_date.toString("dd.MM.yyyy")
                except Exception:
                    fields["birth"] = ""
                self.add_group_visitor_row_custom(fields)
            QMessageBox.information(self, "Успех", "Гости успешно загружены из файла!")
        except ImportError:
            QMessageBox.warning(
                self, "Ошибка", "Установите библиотеку openpyxl: pip install openpyxl"
            )
        except Exception as e:
            QMessageBox.warning(self, "Ошибка", f"Ошибка загрузки файла: {e}")

    def import_visitors(self):
        try:
            from openpyxl import load_workbook
            from PyQt5.QtWidgets import QFileDialog
            from PyQt5.QtCore import QDate
            import re

            file_path, _ = QFileDialog.getOpenFileName(
                self,
                "Выберите файл со списком посетителей",
                "",
                "Excel Files (*.xlsx *.xls)",
            )
            if not file_path:
                return
            wb = load_workbook(file_path)
            ws = wb.active
            headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]
            norm_map = {i: self._normalize_header(h) for i, h in enumerate(headers)}
            required = [
                "last_name",
                "first_name",
                "middle_name",
                "phone",
                "email",
                "org",
                "note",
                "birth",
                "series",
                "number",
            ]
            if not all(v in norm_map.values() for v in required):
                missing = [v for v in required if v not in norm_map.values()]
                QMessageBox.warning(
                    self, "Ошибка", f"В файле отсутствуют столбцы: {', '.join(missing)}"
                )
                return
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            for row in rows:
                if not row:
                    continue
                fields = {}
                for idx, norm in norm_map.items():
                    if norm:
                        fields[norm] = (
                            str(row[idx]).strip() if row[idx] is not None else ""
                        )
                phone_digits = "".join(filter(str.isdigit, fields.get("phone", "")))
                phone_formatted = ""
                if len(phone_digits) == 10:
                    phone_formatted = f"+7 ({phone_digits[0:3]}) {phone_digits[3:6]}-{phone_digits[6:8]}-{phone_digits[8:10]}"
                fields["phone"] = phone_formatted
                try:
                    d, m, y = map(int, str(fields.get("birth", "")).split("."))
                    birth_date = QDate(y, m, d)
                    fields["birth"] = birth_date.toString("dd.MM.yyyy")
                except Exception:
                    fields["birth"] = ""
                self.add_group_visitor_row_custom(fields)
            QMessageBox.information(self, "Успех", "Гости успешно загружены из файла!")
        except ImportError:
            QMessageBox.warning(
                self, "Ошибка", "Установите библиотеку openpyxl: pip install openpyxl"
            )
        except Exception as e:
            QMessageBox.warning(self, "Ошибка", f"Ошибка загрузки файла: {e}")

    def attach_file(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Выберите файл",
            "",
            "Все файлы (*.*);;Документы (*.pdf *.doc *.docx);;Изображения (*.jpg *.jpeg *.png)",
        )

        if file_path:
            self.attach_btn.setText(os.path.basename(file_path))
            self.attach_btn.setProperty("file_path", file_path)


class GuardPanel(QMainWindow):
    def __init__(self, fio):
        super().__init__()
        self.fio = fio
        self.setWindowTitle("Guard")
        self.setMinimumSize(1200, 800)
        self.init_ui()

    def init_ui(self):
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QHBoxLayout(central_widget)  
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        menu_widget = QWidget()
        menu_widget.setStyleSheet(
            """
            background: qlineargradient(x1:0, y1:0, x2:1, y2:1, stop:0 #e0eafc, stop:1 #cfdef3);
            border-top-left-radius: 32px;
            border-bottom-left-radius: 32px;
            box-shadow: 6px 0 32px #b0b0b033;
            """
        )
        menu_layout = QVBoxLayout(menu_widget)
        menu_layout.setAlignment(Qt.AlignTop)
        menu_widget.setFixedWidth(320)

        self.events_btn = QPushButton()
        self.requests_btn = QPushButton()
        self.create_request_btn = QPushButton()

        from PyQt5.QtGui import QPixmap

        avatar = QLabel()
        avatar.setPixmap(QPixmap(64, 64))
        avatar.setStyleSheet(
            "background:#fff;border-radius:32px;border:3px solid #7BA89F;margin-top:32px;margin-bottom:8px;"
        )
        avatar.setFixedSize(64, 64)
        menu_layout.addWidget(avatar, alignment=Qt.AlignHCenter)
        fio_lbl = QLabel(f"<b style='font-size:18px;color:#7BA89F;'>{self.fio}</b>")
        fio_lbl.setAlignment(Qt.AlignCenter)
        fio_lbl.setStyleSheet("margin-bottom:18px;")
        menu_layout.addWidget(fio_lbl)

        title = QLabel(
            "<span style='font-size:28px;font-weight:900;letter-spacing:2px;color:#FF9800;text-shadow:0 2px 8px #7BA89F;'>Вахтер</span>"
        )
        title.setAlignment(Qt.AlignCenter)
        title.setStyleSheet("margin-bottom: 18px; padding: 0 10px;")
        menu_layout.addWidget(title)

        sep = QFrame()
        sep.setFrameShape(QFrame.HLine)
        sep.setStyleSheet("color:#e0e6ed;margin:0 0 18px 0;")
        menu_layout.addWidget(sep)

        btn_style = (
            "QPushButton {"
            "background: #fff; font-size: 17px; padding: 14px 0; margin-bottom: 18px; border-radius: 14px; border: none;"
            "box-shadow: 0 2px 8px #b0b0b022; color: #222; font-weight: 600; letter-spacing: 1px;"
            "transition: background 0.3s, color 0.3s;"
            "text-align: left; padding-left: 32px;"
            "}"
            "QPushButton:hover {"
            "background: #FF9800; color: #fff;"
            "}"
        )
        menu_buttons = [
            (self.events_btn, "📋 Журнал событий"),
            (self.requests_btn, "📄 Заявки"),
            (self.create_request_btn, "➕ Создать заявку"),
        ]
        for btn, text in menu_buttons:
            btn.setText(text)
            btn.setStyleSheet(btn_style)
            btn.setCursor(Qt.PointingHandCursor)
            btn.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
            menu_layout.addWidget(btn)

        self.events_btn.clicked.connect(self.show_events)
        self.requests_btn.clicked.connect(self.show_requests)
        self.create_request_btn.clicked.connect(self.show_create_request_tiles)

        menu_layout.addStretch()
        sep2 = QFrame()
        sep2.setFrameShape(QFrame.HLine)
        sep2.setStyleSheet("color:#e0e6ed;margin:18px 0 0 0;")
        menu_layout.addWidget(sep2)
        version_lbl = QLabel("<span style='color:#bbb;font-size:13px;'>v1.0</span>")
        version_lbl.setAlignment(Qt.AlignCenter)
        version_lbl.setStyleSheet("margin-top:8px;margin-bottom:8px;")
        menu_layout.addWidget(version_lbl)
        logout_btn = QPushButton("Выйти")
        logout_btn.setStyleSheet(
            "background:#eee;font-size:15px;padding:8px 0;border-radius:10px;color:#888;font-weight:600;"
        )
        logout_btn.setCursor(Qt.PointingHandCursor)
        menu_layout.addWidget(logout_btn)
        main_layout.addWidget(menu_widget)
        self.work_area = QWidget()
        self.work_area.setStyleSheet(
            """
            background: #F3F4F6;
            border-top-right-radius: 24px;
            border-bottom-right-radius: 24px;
            """
        )
        self.work_layout = QVBoxLayout(self.work_area)
        main_layout.addWidget(self.work_area)

    def show_events(self):
        self.clear_work_area()
        title = QLabel("Журнал событий")
        title.setStyleSheet(
            """
            font-size: 32px;
            font-weight: bold;
            color: #2c3e50;
            margin-bottom: 20px;
        """
        )
        self.work_layout.addWidget(title)
        events_table = QTableWidget()
        events_table.setStyleSheet(
            """
            QTableWidget {
                background: white;
                border-radius: 16px;
                border: none;
                gridline-color: #E0E6ED;
            }
            QTableWidget::item {
                padding: 8px;
                border-bottom: 1px solid #E0E6ED;
            }
            QHeaderView::section {
                background: #F8F9FA;
                padding: 12px;
                border: none;
                border-bottom: 2px solid #E0E6ED;
                font-weight: bold;
                color: #2c3e50;
            }
        """
        )
        events_table.setColumnCount(3)
        events_table.setHorizontalHeaderLabels(["Время", "Событие", "Статус"])
        events_table.horizontalHeader().setSectionResizeMode(
            0, QHeaderView.ResizeToContents
        )
        events_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        events_table.horizontalHeader().setSectionResizeMode(
            2, QHeaderView.ResizeToContents
        )
        events_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        events_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        events_table.verticalHeader().setVisible(False)
        try:
            with open("traffic_history.json", "r", encoding="utf-8") as f:
                history = json.load(f)
                events_table.setRowCount(len(history))
                for i, event in enumerate(history):
                    time_item = QTableWidgetItem(event.get("time", ""))
                    message_item = QTableWidgetItem(event.get("message", ""))
                    status_item = QTableWidgetItem(
                        "Вход" if "Вход" in event.get("message", "") else "Выход"
                    )
                    status_item.setForeground(
                        QColor(
                            "#27ae60"
                            if "Вход" in event.get("message", "")
                            else "#e74c3c"
                        )
                    )

                    events_table.setItem(i, 0, time_item)
                    events_table.setItem(i, 1, message_item)
                    events_table.setItem(i, 2, status_item)
        except FileNotFoundError:
            pass

        self.work_layout.addWidget(events_table)

    def show_requests(self):
        self.clear_work_area()
        title = QLabel("Заявки на посещение")
        title.setStyleSheet(
            """
            QLabel {
                font-size: 32px;
                font-weight: bold;
                color: #2c3e50;
                margin-bottom: 20px;
            }
        """
        )
        self.work_layout.addWidget(title)

        self.request_tabs = QTabWidget()
        self.request_tabs.setStyleSheet(
            """
            QTabWidget::pane {
                border: 0;
                background: #F3F4F6;
                border-radius: 16px;
                margin-top: 12px;
            }
            QTabBar::tab {
                background: #fff;
                color: #000;
                padding: 12px 24px;
                border-top-left-radius: 12px;
                border-top-right-radius: 12px;
                margin-right: 4px;
                font-size: 16px;
                font-weight: bold;
                min-width: 160px;
            }
            QTabBar::tab:selected {
                background: #FF9800;
                color: #000;
            }
            QTabBar::tab:hover:!selected {
                background: #ffe0b2;
                color: #000;
            }
        """
        )

        all_requests_tab = QWidget()
        all_requests_layout = QVBoxLayout(all_requests_tab)
        self.all_requests_table = self.create_requests_table()
        all_requests_layout.addWidget(self.all_requests_table)
        self.request_tabs.addTab(all_requests_tab, "Все заявки")

        approved_requests_tab = QWidget()
        approved_requests_layout = QVBoxLayout(approved_requests_tab)
        self.approved_requests_table = self.create_requests_table()
        approved_requests_layout.addWidget(self.approved_requests_table)
        self.request_tabs.addTab(approved_requests_tab, "Одобренные")

        rejected_requests_tab = QWidget()
        rejected_requests_layout = QVBoxLayout(rejected_requests_tab)
        self.rejected_requests_table = self.create_requests_table()
        rejected_requests_layout.addWidget(self.rejected_requests_table)
        self.request_tabs.addTab(rejected_requests_tab, "Отклоненные")

        pending_requests_tab = QWidget()
        pending_requests_layout = QVBoxLayout(pending_requests_tab)
        self.pending_requests_table = self.create_requests_table()
        pending_requests_layout.addWidget(self.pending_requests_table)
        self.request_tabs.addTab(pending_requests_tab, "В процессе")

        history_tab = QWidget()
        history_layout = QVBoxLayout(history_tab)

        filter_layout = QHBoxLayout()

        date_filter_label = QLabel("Период:")
        date_filter_label.setStyleSheet("font-size: 14px; color: #666;")
        date_from = QDateEdit()
        date_from.setCalendarPopup(True)
        date_from.setDisplayFormat("dd.MM.yyyy")
        date_from.setDate(QDate.currentDate().addMonths(-1))
        date_to = QDateEdit()
        date_to.setCalendarPopup(True)
        date_to.setDisplayFormat("dd.MM.yyyy")
        date_to.setDate(QDate.currentDate())

        type_filter_label = QLabel("Тип заявки:")
        type_filter_label.setStyleSheet("font-size: 14px; color: #666;")
        type_filter = QComboBox()
        type_filter.addItems(["Все", "Индивидуальная", "Групповая"])

        apply_filter_btn = QPushButton("Применить")
        apply_filter_btn.setStyleSheet(
            """
            QPushButton {
                background-color: #FF9800;
                color: white;
                border: none;
                padding: 8px 16px;
                border-radius: 8px;
                font-weight: bold;
            }
            QPushButton:hover {
                background-color: #F57C00;
            }
        """
        )

        filter_layout.addWidget(date_filter_label)
        filter_layout.addWidget(date_from)
        filter_layout.addWidget(QLabel("—"))
        filter_layout.addWidget(date_to)
        filter_layout.addSpacing(20)
        filter_layout.addWidget(type_filter_label)
        filter_layout.addWidget(type_filter)
        filter_layout.addSpacing(20)
        filter_layout.addWidget(apply_filter_btn)
        filter_layout.addStretch()

        history_layout.addLayout(filter_layout)

        self.history_table = self.create_history_table()
        history_layout.addWidget(self.history_table)

        self.request_tabs.addTab(history_tab, "История")

        self.work_layout.addWidget(self.request_tabs)

        self.update_requests_tables()

        apply_filter_btn.clicked.connect(
            lambda: self.apply_history_filters(
                date_from.date(), date_to.date(), type_filter.currentText()
            )
        )

    def create_history_table(self):
        table = QTableWidget()
        table.setColumnCount(8)
        table.setHorizontalHeaderLabels(
            [
                "№",
                "Тип заявки",
                "ФИО посетителя",
                "Дата создания",
                "Действует с",
                "Действует по",
                "Статус",
                "Действия",
            ]
        )
        table.setMinimumHeight(600)
        table.verticalHeader().setDefaultSectionSize(50)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setSelectionMode(QAbstractItemView.SingleSelection)
        table.setStyleSheet(
            """
            QTableWidget {
                background-color: white;
                border: 1px solid #e0e0e0;
                border-radius: 12px;
                gridline-color: #f0f0f0;
            }
            QTableWidget::item {
                padding: 12px;
                border-bottom: 1px solid #f0f0f0;
            }
            QHeaderView::section {
                background-color: #f8f9fa;
                padding: 16px;
                border: none;
                border-bottom: 2px solid #e0e0e0;
                font-weight: bold;
                color: #2c3e50;
            }
            QTableWidget::item:selected {
                background-color: #fff3e0;
                color: #FF9800;
            }
            QPushButton {
                background-color: #fff;
                color: #FF9800;
                border: 2px solid #FF9800;
                padding: 8px 24px;
                border-radius: 8px;
                font-size: 16px;
                min-width: 120px;
                font-weight: bold;
            }
            QPushButton:hover {
                background-color: #FF9800;
                color: #fff;
            }
        """
        )
        table.setColumnWidth(7, 180)  
        return table

    def apply_history_filters(self, date_from, date_to, request_type):
        requests = self.load_requests()
        filtered_requests = []

        for request in requests:
            created_at = request.get("created_at", "")
            if not created_at:
                continue

            try:
                created_date = QDateTime.fromString(created_at, Qt.ISODate).date()
                if not created_date.isValid():
                    continue

                if date_from <= created_date <= date_to:
                    if (
                        request_type == "Все"
                        or (
                            request_type == "Индивидуальная"
                            and request.get("type") == "individual"
                        )
                        or (
                            request_type == "Групповая"
                            and request.get("type") == "group"
                        )
                    ):
                        filtered_requests.append(request)
            except Exception as e:
                print(f"Ошибка при обработке даты: {e}")
                continue

        self.populate_history_table(filtered_requests)

    def populate_history_table(self, requests_list):
        self.history_table.setRowCount(len(requests_list))
        from PyQt5.QtGui import QIcon

        for row, request in enumerate(requests_list):
            id_item = QTableWidgetItem(str(request.get("id", "")))
            self.history_table.setItem(row, 0, id_item)

            type_text = (
                "Индивидуальная" if request.get("type") == "individual" else "Групповая"
            )
            type_item = QTableWidgetItem(type_text)
            self.history_table.setItem(row, 1, type_item)

            fio = f"{request.get('last_name', '')} {request.get('first_name', '')} {request.get('middle_name', '')}".strip()
            fio_item = QTableWidgetItem(fio)
            self.history_table.setItem(row, 2, fio_item)

            created_item = QTableWidgetItem(request.get("created_at", ""))
            self.history_table.setItem(row, 3, created_item)

            date_from_item = QTableWidgetItem(request.get("date_from", ""))
            self.history_table.setItem(row, 4, date_from_item)

            date_to_item = QTableWidgetItem(request.get("date_to", ""))
            self.history_table.setItem(row, 5, date_to_item)

            status = request.get("status", "pending")
            status_text = {
                "approved": "Одобрена",
                "rejected": "Отклонена",
                "pending": "В процессе",
            }.get(status, "В процессе")

            status_item = QTableWidgetItem(status_text)
            status_item.setForeground(self.get_status_color_brush(status))
            self.history_table.setItem(row, 6, status_item)

            view_btn = QPushButton("Просмотреть")
            view_btn.setStyleSheet(
                "font-size:12px;padding:2px 6px;border-radius:6px;background:#7BA89F;color:#000;"
            )
            view_btn.setProperty("action-btn", True)
            view_btn.setCursor(Qt.PointingHandCursor)
            view_btn.setIcon(QIcon.fromTheme("search") or QIcon("icons/search.png"))
            view_btn.setIconSize(QSize(20, 20))
            view_btn.clicked.connect(
                lambda checked, r=request: self.view_request_details(r)
            )

            cell_widget = QWidget()
            layout = QHBoxLayout(cell_widget)
            layout.addWidget(view_btn)
            layout.setAlignment(Qt.AlignCenter)
            layout.setContentsMargins(0, 0, 0, 0)
            cell_widget.setLayout(layout)
            cell_widget.setStyleSheet("background: #FFF3E0; border-radius: 12px;")
            self.history_table.setCellWidget(row, 7, cell_widget)

    def clear_work_area(self):
        while self.work_layout.count():
            item = self.work_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()

    def show_create_request_tiles(self):
        self.clear_work_area()
        main_container = QWidget()
        main_layout = QVBoxLayout(main_container)
        main_layout.setAlignment(Qt.AlignCenter)
        main_layout.setContentsMargins(50, 50, 50, 50)
        main_layout.setSpacing(40)
        title = QLabel("Создание заявки")
        title.setStyleSheet(
            """
            QLabel {
                font-size: 36px;
                font-weight: bold;
                color: #2c3e50;
                margin-bottom: 40px;
                text-align: center;
            }
        """
        )
        title.setAlignment(Qt.AlignCenter)
        main_layout.addWidget(title)
        tiles_container = QWidget()
        tiles_layout = QHBoxLayout(tiles_container)
        tiles_layout.setSpacing(40)
        tiles_layout.setContentsMargins(0, 0, 0, 0)
        tiles_layout.setAlignment(Qt.AlignCenter)
        individual_tile = QFrame()
        individual_tile.setObjectName("individualTile")
        individual_tile.setStyleSheet(
            """
            QFrame#individualTile {
                background-color: white;
                border-radius: 24px;
                border: 2px solid #e0e0e0;
                min-width: 450px;
                min-height: 550px;
                transition: all 0.3s ease;
            }
            QFrame#individualTile:hover {
                border-color: #FF9800;
                box-shadow: 0 8px 16px rgba(255, 152, 0, 0.15);
                transform: translateY(-5px);
            }
            QLabel {
                color: #2c3e50;
            }
            QPushButton {
                background-color: #FF9800;
                color: white;
                border: none;
                padding: 16px 32px;
                border-radius: 12px;
                font-size: 18px;
                font-weight: bold;
                min-width: 240px;
                transition: all 0.3s ease;
            }
            QPushButton:hover {
                background-color: #F57C00;
                transform: scale(1.05);
            }
        """
        )
        individual_layout = QVBoxLayout(individual_tile)
        individual_layout.setContentsMargins(40, 40, 40, 40)
        individual_layout.setSpacing(30)
        individual_layout.setAlignment(Qt.AlignCenter)

        individual_icon = QLabel()
        individual_icon.setPixmap(
            QPixmap("icons/individual.png").scaled(
                140, 140, Qt.KeepAspectRatio, Qt.SmoothTransformation
            )
        )
        individual_icon.setAlignment(Qt.AlignCenter)
        individual_layout.addWidget(individual_icon)

        individual_title = QLabel("Индивидуальная заявка")
        individual_title.setStyleSheet(
            "font-size: 28px; font-weight: bold; color: #2c3e50;"
        )
        individual_title.setAlignment(Qt.AlignCenter)
        individual_layout.addWidget(individual_title)

        individual_desc = QLabel(
            "Создание заявки для одного посетителя.\n"
            "Включает персональные данные и документы."
        )
        individual_desc.setStyleSheet("font-size: 18px; color: #666; line-height: 1.5;")
        individual_desc.setAlignment(Qt.AlignCenter)
        individual_desc.setWordWrap(True)
        individual_layout.addWidget(individual_desc)

        individual_layout.addStretch()

        individual_btn = QPushButton("Создать заявку")
        individual_btn.setCursor(Qt.PointingHandCursor)
        individual_btn.clicked.connect(self.open_individual_request_window)
        individual_layout.addWidget(individual_btn, alignment=Qt.AlignCenter)

        group_tile = QFrame()
        group_tile.setObjectName("groupTile")
        group_tile.setStyleSheet(
            """
            QFrame#groupTile {
                background-color: white;
                border-radius: 24px;
                border: 2px solid #e0e0e0;
                min-width: 450px;
                min-height: 550px;
                transition: all 0.3s ease;
            }
            QFrame#groupTile:hover {
                border-color: #FF9800;
                box-shadow: 0 8px 16px rgba(255, 152, 0, 0.15);
                transform: translateY(-5px);
            }
            QLabel {
                color: #2c3e50;
            }
            QPushButton {
                background-color: #FF9800;
                color: white;
                border: none;
                padding: 16px 32px;
                border-radius: 12px;
                font-size: 18px;
                font-weight: bold;
                min-width: 240px;
                transition: all 0.3s ease;
            }
            QPushButton:hover {
                background-color: #F57C00;
                transform: scale(1.05);
            }
        """
        )
        group_layout = QVBoxLayout(group_tile)
        group_layout.setContentsMargins(40, 40, 40, 40)
        group_layout.setSpacing(30)
        group_layout.setAlignment(Qt.AlignCenter)

        group_icon = QLabel()
        group_icon.setPixmap(
            QPixmap("icons/group.png").scaled(
                140, 140, Qt.KeepAspectRatio, Qt.SmoothTransformation
            )
        )
        group_icon.setAlignment(Qt.AlignCenter)
        group_layout.addWidget(group_icon)

        group_title = QLabel("Групповая заявка")
        group_title.setStyleSheet("font-size: 28px; font-weight: bold; color: #2c3e50;")
        group_title.setAlignment(Qt.AlignCenter)
        group_layout.addWidget(group_title)

        group_desc = QLabel(
            "Создание заявки для группы посетителей.\n"
            "Возможность загрузки списка из файла."
        )
        group_desc.setStyleSheet("font-size: 18px; color: #666; line-height: 1.5;")
        group_desc.setAlignment(Qt.AlignCenter)
        group_desc.setWordWrap(True)
        group_layout.addWidget(group_desc)

        group_layout.addStretch()

        group_btn = QPushButton("Создать заявку")
        group_btn.setCursor(Qt.PointingHandCursor)
        group_btn.clicked.connect(self.open_group_request_window)
        group_layout.addWidget(group_btn, alignment=Qt.AlignCenter)

        tiles_layout.addWidget(individual_tile)
        tiles_layout.addWidget(group_tile)

        main_layout.addWidget(tiles_container)
        main_layout.addStretch()

        self.work_layout.addWidget(main_container)
        self.work_layout.addStretch()

    def open_individual_request_window(self):
        dialog = CreateIndividualRequestWindow(self)
        dialog.exec_()

    def open_group_request_window(self):
        dialog = CreateGroupRequestWindow(self)
        dialog.exec_()

    def create_requests_table(self):

        table = QTableWidget()
        table.setColumnCount(7)
        table.setHorizontalHeaderLabels(
            [
                "№",
                "Тип заявки",
                "ФИО посетителя",
                "Дата создания",
                "Действует с",
                "Действует по",
                "Статус",
            ]
        )
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setSelectionMode(QAbstractItemView.SingleSelection)
        table.setStyleSheet(
            """
            QTableWidget {
                background-color: white;
                border: 1px solid #e0e0e0;
                border-radius: 12px;
                gridline-color: #f0f0f0;
            }
            QTableWidget::item {
                padding: 12px;
                border-bottom: 1px solid #f0f0f0;
            }
            QHeaderView::section {
                background-color: #f8f9fa;
                padding: 16px;
                border: none;
                border-bottom: 2px solid #e0e0e0;
                font-weight: bold;
                color: #2c3e50;
            }
            QTableWidget::item:selected {
                background-color: #fff3e0;
                color: #FF9800;
            }
        """
        )
        return table

    def update_requests_tables(self):
        requests = self.load_requests()

        approved_requests = [r for r in requests if r.get("status") == "approved"]
        rejected_requests = [r for r in requests if r.get("status") == "rejected"]
        pending_requests = [r for r in requests if r.get("status") == "pending"]

        self.populate_requests_table(self.all_requests_table, requests)
        self.populate_requests_table(self.approved_requests_table, approved_requests)
        self.populate_requests_table(self.rejected_requests_table, rejected_requests)
        self.populate_requests_table(self.pending_requests_table, pending_requests)

    def populate_requests_table(self, table, requests_list):
        table.setRowCount(len(requests_list))

        for row, request in enumerate(requests_list):
            id_item = QTableWidgetItem(str(request.get("id", "")))
            table.setItem(row, 0, id_item)

            type_text = (
                "Индивидуальная" if request.get("type") == "individual" else "Групповая"
            )
            type_item = QTableWidgetItem(type_text)
            table.setItem(row, 1, type_item)

            if request.get("type") == "individual":
                fio = f"{request.get('last_name', '')} {request.get('first_name', '')} {request.get('middle_name', '')}".strip()
            else:
                visitors = request.get("visitors", [])
                if visitors:
                    fio = visitors[0].get("fio", "")
                else:
                    fio = "Нет данных"
            fio_item = QTableWidgetItem(fio)
            table.setItem(row, 2, fio_item)
            created_at = request.get("created_at", "")
            created_item = QTableWidgetItem(created_at)
            table.setItem(row, 3, created_item)
            date_from = request.get("date_from", "")
            date_from_item = QTableWidgetItem(date_from)
            table.setItem(row, 4, date_from_item)
            date_to = request.get("date_to", "")
            date_to_item = QTableWidgetItem(date_to)
            table.setItem(row, 5, date_to_item)
            status = request.get("status", "pending")
            status_text = {
                "approved": "Одобрена",
                "rejected": "Отклонена",
                "pending": "В процессе",
            }.get(status, "В процессе")
            status_item = QTableWidgetItem(status_text)
            status_item.setForeground(self.get_status_color_brush(status))
            table.setItem(row, 6, status_item)
            if request.get("type") == "group":
                guests_count = QTableWidgetItem(str(len(request.get("visitors", []))))
            else:
                guests_count = QTableWidgetItem("1")
            table.setItem(row, 7, guests_count)

        table.resizeColumnsToContents()
        table.horizontalHeader().setStretchLastSection(True)

    def get_status_color_brush(self, status):
        from PyQt5.QtGui import QColor

        colors = {
            "approved": QColor("#4CAF50"),  
            "rejected": QColor("#F44336"),  
            "pending": QColor("#FF9800"),  
        }
        return colors.get(status, QColor("#FF9800"))

    def load_requests(self):
        try:
            if os.path.exists(REQUESTS_FILE):
                with open(REQUESTS_FILE, "r", encoding="utf-8") as f:
                    return json.load(f)
            return []
        except Exception as e:
            print(f"Ошибка при загрузке заявок: {e}")
            return []

    def save_requests(self, data):
        try:
            with open(REQUESTS_FILE, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print(f"Ошибка при сохранении заявок: {e}")

    def view_request_details(self, request):
        if request.get("type") == "individual":
            dialog = ViewEditIndividualRequestDialog(
                request, is_editable=False, parent=self
            )
        else:
            dialog = ViewEditGroupRequestDialog(request, is_editable=False, parent=self)
        dialog.exec_()

    def _normalize_header(self, header):
        header = (
            header.lower()
            .replace("ё", "е")
            .replace(" ", "")
            .replace("-", "")
            .replace("(", "")
            .replace(")", "")
            .replace(".", "")
        )
        if "фамил" in header:
            return "last_name"
        if "имя" in header:
            return "first_name"
        if "отчест" in header:
            return "middle_name"
        if "телефон" in header:
            return "phone"
        if "email" in header:
            return "email"
        if "организац" in header:
            return "org"
        if "примечан" in header:
            return "note"
        if "датарожд" in header:
            return "birth"
        if "серия" in header:
            return "series"
        if "номер" in header:
            return "number"
        return None


class PhoneValidator(QRegExpValidator):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setRegExp(QRegExp(r"\+7 \(\d{3}\) \d{3}-\d{2}-\d{2}"))

    def validate(self, string, pos):
        if not string:
            return (QValidator.Acceptable, string, pos)

        digits = "".join(filter(str.isdigit, string))

        if len(digits) > 11:
            return (QValidator.Invalid, string, pos)
        if len(digits) >= 1:
            string = "+7 ("
        if len(digits) >= 4:
            string += f"{digits[1:4]}) "
        if len(digits) >= 7:
            string += f"{digits[4:7]}-"
        if len(digits) >= 9:
            string += f"{digits[7:9]}-"
        if len(digits) >= 11:
            string += f"{digits[9:11]}"

        return (QValidator.Acceptable, string, pos)


def validate_email(email):
    pattern = r"^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$"
    return bool(re.match(pattern, email))


def validate_passport(series, number):
    series = series.replace(" ", "")
    number = number.replace(" ", "")
    return (
        len(series) == 4 and series.isdigit() and len(number) == 6 and number.isdigit()
    )


def validate_photo(file_path):
    if not file_path:
        return True, ""

    try:
        from PIL import Image

        img = Image.open(file_path)
        width, height = img.size
        file_size = os.path.getsize(file_path) / (1024 * 1024)
        if file_size > 4:
            return False, "Размер файла превышает 4 МБ"

        if img.format.lower() not in ["jpeg", "jpg", "png"]:
            return False, "Формат файла должен быть JPG или PNG"



        return True, ""
    except Exception as e:
        return False, f"Ошибка при проверке фото: {str(e)}"


def validate_age(birth_date, min_age=14):
    from PyQt5.QtCore import QDate

    today = QDate.currentDate()

    if not birth_date or not isinstance(birth_date, QDate) or not birth_date.isValid():
        return False
    age = today.year() - birth_date.year()
    if birth_date > today.addYears(-age):
        age -= 1
    return age >= min_age


class PassportValidator(QValidator):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.regex = QRegExp(r"^\d{4} \d{6}$")

    def validate(self, string, pos):
        if not string:
            return (QValidator.Acceptable, string, pos)

        digits = "".join(filter(str.isdigit, string))

        if len(digits) <= 4:
            return (QValidator.Acceptable, digits, pos)

        formatted = f"{digits[:4]} {digits[4:10]}"
        return (QValidator.Acceptable, formatted, pos)

    def fixup(self, string):
        digits = "".join(filter(str.isdigit, string))
        if len(digits) >= 4:
            return f"{digits[:4]} {digits[4:10]}"
        return string


class DummyField:
    def __init__(self, value):
        self._value = value if value is not None else ""

    def text(self):
        return str(self._value)


class DummyDateField:
    def __init__(self, value):
        from PyQt5.QtCore import QDate

        if value:
            try:
                d, m, y = map(int, str(value).split("."))
                self._date = QDate(y, m, d)
            except:
                self._date = QDate.currentDate()
        else:
            self._date = QDate.currentDate()

    def date(self):
        return self._date

    def is_valid(self):
        return validate_age(self._date, min_age=16)

    def text(self):

        return self._date.toString("dd.MM.yyyy")


class CreateIndividualRequestWindow(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Создание индивидуальной заявки")
        self.setMinimumSize(1200, 800)
        self.init_ui()

    def init_ui(self):
        main_layout = QVBoxLayout(self)

        self.setStyleSheet(
            "background: qlineargradient(x1:0, y1:0, x2:1, y2:1, stop:0 #f8fafc, stop:1 #e0eafc);"
        )

        card = QFrame()
        card.setStyleSheet(
            """
            background: #fff;
            border-radius: 32px;
            box-shadow: 0 8px 32px rgba(0, 0, 0, 0.1);
            padding: 0px;
            width: 100%;
            margin-top: 24px;
            margin-bottom: 24px;
            """
        )
        card_layout = QVBoxLayout(card)
        card_layout.setContentsMargins(40, 40, 40, 40)
        card_layout.setSpacing(30)

        title_box = QHBoxLayout()
        icon = QLabel()
        icon.setPixmap(QIcon("icons/individual.png").pixmap(48, 48))
        icon.setFixedSize(56, 56)
        title_box.addWidget(icon)
        title = QLabel(
            "<span style='font-size:32px;color:#FF9800;font-weight:bold;'>Индивидуальная заявка</span>"
        )
        title.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        title_box.addWidget(title)
        title_box.addStretch()
        card_layout.addLayout(title_box)

        form_layout = QVBoxLayout()
        form_layout.setSpacing(24)

        pass_frame = QFrame()
        pass_frame.setStyleSheet(
            """
            background: #F8F9FA;
            border-radius: 24px;
            border: 2px solid #E0E6ED;
            """
        )
        pass_layout = QVBoxLayout(pass_frame)
        pass_title = QLabel("Информация о пропуске")
        pass_title.setProperty("role", "block-title")
        pass_title.setStyleSheet(
            """
            background: #FF9800;
            color: #fff;
            font-size: 22px;
            font-weight: bold;
            border-radius: 16px;
            padding: 12px 36px;
            margin-bottom: 24px;
            """
        )
        pass_layout.addWidget(pass_title)

        pass_grid = QGridLayout()
        pass_grid.setSpacing(20)

        self.pass_id = QLineEdit()
        self.pass_id.setText(str(random.randint(100000, 999999)))
        self.pass_id.setReadOnly(True)
        self.pass_id.setStyleSheet(
            """
            QLineEdit {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            QLineEdit:read-only {
                background: #F8F9FA;
            }
            """
        )
        pass_grid.addWidget(QLabel("Номер пропуска:"), 0, 0)
        pass_grid.addWidget(self.pass_id, 0, 1)
        self.temp_pass = QCheckBox("Временный пропуск")
        self.temp_pass.setStyleSheet(
            """
            QCheckBox {
                font-size: 16px;
                padding: 12px;
            }
            QCheckBox::indicator {
                width: 20px;
                height: 20px;
            }
        """
        )
        pass_grid.addWidget(self.temp_pass, 0, 2, 1, 2)

        self.date_from = QDateEdit()
        self.date_from.setCalendarPopup(True)
        self.date_from.setDisplayFormat("dd.MM.yyyy")
        self.date_from.setDate(QDate.currentDate())
        self.date_from.setStyleSheet(
            """
            QDateEdit {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            QDateEdit::drop-down {
                border: none;
                width: 30px;
            }
            """
        )

        self.date_to = QDateEdit()
        self.date_to.setCalendarPopup(True)
        self.date_to.setDisplayFormat("dd.MM.yyyy")
        self.date_to.setDate(QDate.currentDate())
        self.date_to.setStyleSheet(
            """
            QDateEdit {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            QDateEdit::drop-down {
                border: none;
                width: 30px;
            }
            """
        )

        pass_grid.addWidget(QLabel("Действует с:"), 1, 0)
        pass_grid.addWidget(self.date_from, 1, 1)
        pass_grid.addWidget(QLabel("по:"), 1, 2)
        pass_grid.addWidget(self.date_to, 1, 3)

        self.time_from = QTimeEdit()
        self.time_from.setDisplayFormat("HH:mm")
        self.time_from.setTime(QTime(9, 0))
        self.time_from.setStyleSheet(
            """
            QTimeEdit {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            QTimeEdit::drop-down {
                border: none;
                width: 30px;
            }
        """
        )

        self.time_to = QTimeEdit()
        self.time_to.setDisplayFormat("HH:mm")
        self.time_to.setTime(QTime(18, 0))
        self.time_to.setStyleSheet(
            """
            QTimeEdit {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            QTimeEdit::drop-down {
                border: none;
                width: 30px;
            }
        """
        )

        pass_grid.addWidget(QLabel("Время с:"), 2, 0)
        pass_grid.addWidget(self.time_from, 2, 1)
        pass_grid.addWidget(QLabel("до:"), 2, 2)
        pass_grid.addWidget(self.time_to, 2, 3)

        self.purpose = QComboBox()
        self.purpose.addItems(["Деловая встреча", "Собеседование", "Другое"])
        self.purpose.setStyleSheet(
            """
            QComboBox {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            QComboBox::drop-down {
                border: none;
                width: 30px;
            }
            """
        )
        self.purpose.currentTextChanged.connect(self.on_purpose_changed)
        self.purpose_other = QLineEdit()
        self.purpose_other.setPlaceholderText("Укажите цель посещения")
        self.purpose_other.setStyleSheet(
            """
            QLineEdit {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            """
        )
        self.purpose_other.hide()
        pass_grid.addWidget(QLabel("Цель посещения:"), 3, 0)
        pass_grid.addWidget(self.purpose, 3, 1, 1, 3)
        pass_grid.addWidget(self.purpose_other, 4, 1, 1, 3)
        pass_layout.addLayout(pass_grid)
        form_layout.addWidget(pass_frame)

        host_frame = QFrame()
        host_frame.setStyleSheet(
            """
            background: #F8F9FA;
            border-radius: 24px;
            border: 2px solid #E0E6ED;
            """
        )
        host_layout = QVBoxLayout(host_frame)
        host_title = QLabel("Принимающая сторона")
        host_title.setProperty("role", "block-title")
        host_title.setStyleSheet(
            """
            background: #FF9800;
            color: #fff;
            font-size: 22px;
            font-weight: bold;
            border-radius: 16px;
            padding: 12px 36px;
            margin-bottom: 24px;
            """
        )
        host_layout.addWidget(host_title)

        host_grid = QGridLayout()
        host_grid.setSpacing(20)

        self.department = QLineEdit()
        self.department.setPlaceholderText("Подразделение")
        self.department.setStyleSheet(
            """
            QLineEdit {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            """
        )
        self.host_fio = QLineEdit()
        self.host_fio.setPlaceholderText("ФИО принимающей стороны")
        self.host_fio.setStyleSheet(
            """
            QLineEdit {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            """
        )

        host_grid.addWidget(QLabel("Подразделение:"), 0, 0)
        host_grid.addWidget(self.department, 0, 1)
        host_grid.addWidget(QLabel("ФИО принимающей стороны:"), 1, 0)
        host_grid.addWidget(self.host_fio, 1, 1)
        host_layout.addLayout(host_grid)
        form_layout.addWidget(host_frame)

        visitor_frame = QFrame()
        visitor_frame.setStyleSheet(
            """
            background: #F8F9FA;
            border-radius: 24px;
            border: 2px solid #E0E6ED;
            """
        )
        visitor_layout = QVBoxLayout(visitor_frame)
        visitor_title = QLabel("Данные посетителя")
        visitor_title.setProperty("role", "block-title")
        visitor_title.setStyleSheet(
            """
            background: #FF9800;
            color: #fff;
            font-size: 22px;
            font-weight: bold;
            border-radius: 16px;
            padding: 12px 36px;
            margin-bottom: 24px;
            """
        )
        visitor_layout.addWidget(visitor_title)

        visitor_grid = QGridLayout()
        visitor_grid.setSpacing(20)

        self.last_name = QLineEdit()
        self.last_name.setPlaceholderText("Фамилия")
        self.first_name = QLineEdit()
        self.first_name.setPlaceholderText("Имя")
        self.middle_name = QLineEdit()
        self.middle_name.setPlaceholderText("Отчество")

        for field in [self.last_name, self.first_name, self.middle_name]:
            field.setStyleSheet(
                """
                QLineEdit {
                    padding: 12px;
                    border: 2px solid #E0E6ED;
                    border-radius: 12px;
                    font-size: 16px;
                    background: white;
                }
                """
            )

        visitor_grid.addWidget(QLabel("Фамилия:"), 0, 0)
        visitor_grid.addWidget(self.last_name, 0, 1)
        visitor_grid.addWidget(QLabel("Имя:"), 0, 2)
        visitor_grid.addWidget(self.first_name, 0, 3)
        visitor_grid.addWidget(QLabel("Отчество:"), 1, 0)
        visitor_grid.addWidget(self.middle_name, 1, 1)

        self.phone = QLineEdit()
        self.phone.setPlaceholderText("+7 (XXX) XXX-XX-XX")
        self.phone.textChanged.connect(self.format_phone)
        self.email = QLineEdit()
        self.email.setPlaceholderText("Email")
        self.org = QLineEdit()
        self.org.setPlaceholderText("Организация")

        for field in [self.phone, self.email, self.org]:
            field.setStyleSheet(
                """
                QLineEdit {
                    padding: 12px;
                    border: 2px solid #E0E6ED;
                    border-radius: 12px;
                    font-size: 16px;
                    background: white;
                }
                """
            )

        visitor_grid.addWidget(QLabel("Телефон:"), 2, 0)
        visitor_grid.addWidget(self.phone, 2, 1)
        visitor_grid.addWidget(QLabel("Email:"), 2, 2)
        visitor_grid.addWidget(self.email, 2, 3)
        visitor_grid.addWidget(QLabel("Организация:"), 3, 0)
        visitor_grid.addWidget(self.org, 3, 1)

        self.birth = QDateEdit()
        self.birth.setCalendarPopup(True)
        self.birth.setDisplayFormat("dd.MM.yyyy")
        self.birth.setDate(QDate.currentDate())
        self.series = QLineEdit()
        self.series.setPlaceholderText("Серия")
        self.series.setValidator(PassportValidator())
        self.number = QLineEdit()
        self.number.setPlaceholderText("Номер")
        self.number.setValidator(PassportValidator())

        for field in [self.birth, self.series, self.number]:
            field.setStyleSheet(
                """
                QLineEdit, QDateEdit {
                    padding: 12px;
                    border: 2px solid #E0E6ED;
                    border-radius: 12px;
                    font-size: 16px;
                    background: white;
                }
                QDateEdit::drop-down {
                    border: none;
                    width: 30px;
                }
                """
            )

        visitor_grid.addWidget(QLabel("Дата рождения:"), 4, 0)
        visitor_grid.addWidget(self.birth, 4, 1)
        visitor_grid.addWidget(QLabel("Паспорт:"), 4, 2)
        visitor_grid.addWidget(self.series, 4, 3)
        visitor_grid.addWidget(self.number, 4, 4)

        self.note = QTextEdit()
        self.note.setPlaceholderText("Примечание")
        self.note.setStyleSheet(
            """
            QTextEdit {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
                min-height: 100px;
            }
            """
        )
        visitor_grid.addWidget(QLabel("Примечание:"), 5, 0)
        visitor_grid.addWidget(self.note, 5, 1, 1, 4)

        visitor_layout.addLayout(visitor_grid)
        form_layout.addWidget(visitor_frame)

        docs_frame = QFrame()
        docs_frame.setStyleSheet(
            """
            background: #F8F9FA;
            border-radius: 24px;
            border: 2px solid #E0E6ED;
            """
        )
        docs_layout = QVBoxLayout(docs_frame)
        docs_title = QLabel("Документы")
        docs_title.setProperty("role", "block-title")
        docs_title.setStyleSheet(
            """
            background: #FF9800;
            color: #fff;
            font-size: 22px;
            font-weight: bold;
            border-radius: 16px;
            padding: 12px 36px;
            margin-bottom: 24px;
            """
        )
        docs_layout.addWidget(docs_title)

        docs_grid = QGridLayout()
        docs_grid.setSpacing(20)

        self.photo_label = QLabel("Фото")
        self.photo_label.setAlignment(Qt.AlignCenter)
        self.photo_label.setStyleSheet(
            """
            QLabel {
                padding: 20px;
                border: 2px dashed #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                color: #666;
                background: white;
                min-height: 200px;
            }
            """
        )
        self.photo_label.setCursor(Qt.PointingHandCursor)
        self.photo_label.mousePressEvent = self.load_photo

        self.attach_btn = QPushButton("Прикрепить файл")
        self.attach_btn.setStyleSheet(
            """
            QPushButton {
                padding: 12px 24px;
                border: 2px solid #FF9800;
                border-radius: 12px;
                font-size: 16px;
                font-weight: bold;
                color: #FF9800;
                background: white;
            }
            QPushButton:hover {
                background: #FF9800;
                color: white;
            }
            """
        )
        self.attach_btn.setCursor(Qt.PointingHandCursor)
        self.attach_btn.clicked.connect(self.attach_file)

        docs_grid.addWidget(QLabel("Фото:"), 0, 0)
        docs_grid.addWidget(self.photo_label, 0, 1)
        docs_grid.addWidget(QLabel("Дополнительный файл:"), 1, 0)
        docs_grid.addWidget(self.attach_btn, 1, 1)

        docs_layout.addLayout(docs_grid)
        form_layout.addWidget(docs_frame)

        card_layout.addLayout(form_layout)

        btns = QHBoxLayout()
        clear_btn = QPushButton("Очистить форму")
        clear_btn.setStyleSheet(
            """
            QPushButton {
                padding: 16px 32px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 18px;
                font-weight: bold;
                color: #666;
                background: white;
            }
            QPushButton:hover {
                background: #F8F9FA;
            }
            """
        )
        clear_btn.setCursor(Qt.PointingHandCursor)
        clear_btn.clicked.connect(self.clear_individual_form)

        submit_btn = QPushButton("Оформить заявку")
        submit_btn.setStyleSheet(
            """
            QPushButton {
                padding: 16px 32px;
                border: none;
                border-radius: 12px;
                font-size: 18px;
                font-weight: bold;
                color: white;
                background: #FF9800;
            }
            QPushButton:hover {
                background: #F57C00;
            }
            """
        )
        submit_btn.setCursor(Qt.PointingHandCursor)
        submit_btn.clicked.connect(self.submit_individual_form)

        btns.addWidget(clear_btn)
        btns.addWidget(submit_btn)
        card_layout.addLayout(btns)
        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)
        scroll_area.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        scroll_area.setWidget(card)
        scroll_area.setStyleSheet("border: none; background: transparent;")
        main_layout.addWidget(scroll_area)

    def load_photo(self, *args):
        fname, _ = QFileDialog.getOpenFileName(
            self, "Выбрать фото", "", "Изображения (*.png *.jpg *.jpeg)"
        )
        if fname:
            try:
                from PIL import Image
                from PyQt5.QtGui import QPixmap
                import io

                img = Image.open(fname)
                width, height = img.size
                target_ratio = 3 / 4
                img_ratio = width / height
                if abs(img_ratio - target_ratio) > 0.01:
                    if img_ratio > target_ratio:
                        new_width = int(height * target_ratio)
                        left = (width - new_width) // 2
                        img = img.crop((left, 0, left + new_width, height))
                    else:
                        new_height = int(width / target_ratio)
                        top = (height - new_height) // 2
                        img = img.crop((0, top, width, top + new_height))

                img = img.resize(
                    (self.photo_label.width(), self.photo_label.height()), Image.LANCZOS
                )
                buf = io.BytesIO()
                img.save(buf, format="PNG")
                qt_pixmap = QPixmap()
                qt_pixmap.loadFromData(buf.getvalue(), "PNG")
                self.photo_label.setPixmap(qt_pixmap)
                self.photo_label.setProperty("file_path", fname)
                self.photo_label.setText("")
            except Exception as e:
                QMessageBox.warning(self, "Ошибка", f"Ошибка обработки фото: {e}")
                self.photo_label.setText("Фото")
                self.photo_label.setProperty("file_path", "")
                self.photo_label.clear()
        else:
            self.photo_label.setText("Фото")
            self.photo_label.setProperty("file_path", "")
            self.photo_label.clear()

    def attach_file(self):
        file_name, _ = QFileDialog.getOpenFileName(
            self, "Выберите файл", "", "Все файлы (*.*)"
        )
        if file_name:
            self.attach_btn.setText(os.path.basename(file_name))
            self.attach_btn.setProperty("file_path", file_name)
        else:
            self.attach_btn.setText("Прикрепить файл")
            self.attach_btn.setProperty("file_path", "")

    def clear_individual_form(self):
        self.pass_id.setText(str(random.randint(100000, 999999)))
        self.temp_pass.setChecked(False)
        self.time_from.setTime(QTime(9, 0))
        self.time_to.setTime(QTime(18, 0))
        self.date_from.setDate(QDate.currentDate())
        self.date_to.setDate(QDate.currentDate())
        self.purpose.setCurrentIndex(0)
        self.purpose_other.clear()
        self.purpose_other.hide()
        self.department.clear()
        self.host_fio.clear()
        self.last_name.clear()
        self.first_name.clear()
        self.middle_name.clear()
        self.phone.clear()
        self.email.clear()
        self.org.clear()
        self.note.clear()
        self.birth.setDate(QDate.currentDate())
        self.series.clear()
        self.number.clear()
        self.photo_label.clear()
        self.photo_label.setText("Фото")
        self.photo_label.setProperty("file_path", "")
        self.attach_btn.setText("Прикрепить файл")
        self.attach_btn.setProperty("file_path", "")

    def submit_individual_form(self):
        errors = []

        if not self.department.text().strip():
            errors.append("Подразделение обязательно")
        if not self.host_fio.text().strip():
            errors.append("ФИО принимающей стороны обязательно")

        if not self.last_name.text().strip():
            errors.append("Фамилия посетителя обязательна")
        if not self.first_name.text().strip():
            errors.append("Имя посетителя обязательно")

        phone = self.phone.text().strip()
        if phone and not re.match(r"\+7 \(\d{3}\) \d{3}-\d{2}-\d{2}", phone):
            errors.append("Телефон должен быть в формате +7 (###) ###-##-##")

        email = self.email.text().strip()
        if not email:
            errors.append("Email посетителя обязателен")
        elif not validate_email(email):
            errors.append("Неверный формат Email")

        if not self.note.toPlainText().strip():
            errors.append("Примечание обязательно")

        birth_date = self.birth.date()
        if not validate_age(birth_date):
            errors.append("Возраст посетителя должен быть не моложе 14 лет")

        series = self.series.text().strip()
        number = self.number.text().strip()
        if not validate_passport(series, number):
            errors.append("Серия паспорта должна содержать 4 цифры, номер - 6 цифр")

        photo_path = self.photo_label.property("file_path")
        if photo_path:
            is_valid, error_msg = validate_photo(photo_path)
            if not is_valid:
                errors.append(error_msg)

        attached_file_path = self.attach_btn.property("file_path")
        if not attached_file_path:
            errors.append("Скан паспорта обязателен")
        elif not attached_file_path.lower().endswith(".jpg"):
            errors.append("Скан паспорта должен быть в формате JPG")

        if errors:
            QMessageBox.warning(self, "Ошибка валидации", "\n".join(errors))
            return
        request_id = str(uuid.uuid4())
        created_at = QDateTime.currentDateTime().toString(Qt.ISODate)

        req = {
            "id": request_id,
            "type": "individual",
            "status": "pending",
            "created_at": created_at,
            "pass_id": self.pass_id.text(),
            "date_from": self.date_from.date().toString("yyyy-MM-dd"),
            "date_to": self.date_to.date().toString("yyyy-MM-dd"),
            "purpose": (
                self.purpose.currentText()
                if self.purpose.currentText() != "Другое"
                else self.purpose_other.text()
            ),
            "department": self.department.text(),
            "host_fio": self.host_fio.text(),
            "last_name": self.last_name.text(),
            "first_name": self.first_name.text(),
            "middle_name": self.middle_name.text(),
            "phone": self.phone.text(),
            "email": self.email.text(),
            "org": self.org.text(),
            "note": self.note.toPlainText(),
            "birth": self.birth.date().toString("yyyy-MM-dd"),
            "series": self.series.text(),
            "number": self.number.text(),
            "photo_path": (
                self.photo_label.property("file_path")
                if self.photo_label.property("file_path")
                else ""
            ),
            "attached_file_path": (
                self.attach_btn.property("file_path")
                if self.attach_btn.property("file_path")
                else ""
            ),
        }
        data = self.parent().load_requests()
        data.append(req)
        self.parent().save_requests(data)

        QMessageBox.information(
            self,
            "Успех",
            f"Заявка №{request_id} успешно отправлена администратору на одобрение!",
        )
        self.clear_individual_form()
        self.accept()

    def on_purpose_changed(self, text):
        self.purpose_other.setVisible(text == "Другое")
        if text != "Другое":
            self.purpose_other.clear()

    def format_phone(self, text):
        digits = "".join(filter(str.isdigit, text))

        if len(digits) > 11:
            digits = digits[:11]

        formatted = ""
        if len(digits) > 0:
            formatted = "+7"
        if len(digits) > 1:
            formatted += f" ({digits[1:4]}"
        if len(digits) > 4:
            formatted += f") {digits[4:7]}"
        if len(digits) > 7:
            formatted += f"-{digits[7:9]}"
        if len(digits) > 9:
            formatted += f"-{digits[9:11]}"

        self.phone.blockSignals(True)
        self.phone.setText(formatted)
        self.phone.blockSignals(False)


class CreateGroupRequestWindow(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Создание групповой заявки")
        self.setMinimumSize(1200, 800)
        self.init_ui()

    def init_ui(self):
        main_layout = QVBoxLayout(self)

        self.setStyleSheet(
            "background: qlineargradient(x1:0, y1:0, x2:1, y2:1, stop:0 #f8fafc, stop:1 #e0eafc);"
        )

        card = QFrame()
        card.setStyleSheet(
            """
            background: #fff;
            border-radius: 32px;
            box-shadow: 0 8px 32px rgba(0, 0, 0, 0.1);
            padding: 0px;
            width: 100%;
            margin-top: 24px;
            margin-bottom: 24px;
            """
        )
        card_layout = QVBoxLayout(card)
        card_layout.setContentsMargins(40, 40, 40, 40)
        card_layout.setSpacing(30)

        title_box = QHBoxLayout()
        icon = QLabel()
        icon.setPixmap(QIcon("icons/group.png").pixmap(48, 48))
        icon.setFixedSize(56, 56)
        title_box.addWidget(icon)
        title = QLabel(
            "<span style='font-size:32px;color:#FF9800;font-weight:bold;'>Групповая заявка</span>"
        )
        title.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        title_box.addWidget(title)
        title_box.addStretch()
        card_layout.addLayout(title_box)

        form_layout = QVBoxLayout()
        form_layout.setSpacing(24)

        pass_frame = QFrame()
        pass_frame.setStyleSheet(
            """
            background: #F8F9FA;
            border-radius: 24px;
            border: 2px solid #E0E6ED;
            """
        )
        pass_layout = QVBoxLayout(pass_frame)
        pass_title = QLabel("Информация о пропуске")
        pass_title.setProperty("role", "block-title")
        pass_title.setStyleSheet(
            """
            background: #FF9800;
            color: #fff;
            font-size: 22px;
            font-weight: bold;
            border-radius: 16px;
            padding: 12px 36px;
            margin-bottom: 24px;
            """
        )
        pass_layout.addWidget(pass_title)

        pass_grid = QGridLayout()
        pass_grid.setSpacing(20)

        self.group_pass_id = QLineEdit()
        self.group_pass_id.setText(str(random.randint(100000, 999999)))
        self.group_pass_id.setReadOnly(True)
        self.group_pass_id.setStyleSheet(
            """
            QLineEdit {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            QLineEdit:read-only {
                background: #F8F9FA;
            }
            """
        )
        pass_grid.addWidget(QLabel("Номер пропуска:"), 0, 0)
        pass_grid.addWidget(self.group_pass_id, 0, 1)

        self.group_date_from = QDateEdit()
        self.group_date_from.setCalendarPopup(True)
        self.group_date_from.setDisplayFormat("dd.MM.yyyy")
        self.group_date_from.setDate(QDate.currentDate())
        self.group_date_from.setStyleSheet(
            """
            QDateEdit {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            QDateEdit::drop-down {
                border: none;
                width: 30px;
            }
            """
        )

        self.group_date_to = QDateEdit()
        self.group_date_to.setCalendarPopup(True)
        self.group_date_to.setDisplayFormat("dd.MM.yyyy")
        self.group_date_to.setDate(QDate.currentDate())
        self.group_date_to.setStyleSheet(
            """
            QDateEdit {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            QDateEdit::drop-down {
                border: none;
                width: 30px;
            }
            """
        )

        pass_grid.addWidget(QLabel("Действует с:"), 1, 0)
        pass_grid.addWidget(self.group_date_from, 1, 1)
        pass_grid.addWidget(QLabel("по:"), 1, 2)
        pass_grid.addWidget(self.group_date_to, 1, 3)
        self.group_purpose = QComboBox()
        self.group_purpose.addItems(["Деловая встреча", "Собеседование", "Другое"])
        self.group_purpose.setStyleSheet(
            """
            QComboBox {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            QComboBox::drop-down {
                border: none;
                width: 30px;
            }
            """
        )
        self.group_purpose.currentTextChanged.connect(self.on_purpose_changed)
        self.group_purpose_other = QLineEdit()
        self.group_purpose_other.setPlaceholderText("Укажите цель посещения")
        self.group_purpose_other.setStyleSheet(
            """
            QLineEdit {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            """
        )
        self.group_purpose_other.hide()
        pass_grid.addWidget(QLabel("Цель посещения:"), 2, 0)
        pass_grid.addWidget(self.group_purpose, 2, 1, 1, 3)
        pass_grid.addWidget(self.group_purpose_other, 3, 1, 1, 3)
        pass_layout.addLayout(pass_grid)
        form_layout.addWidget(pass_frame)

        host_frame = QFrame()
        host_frame.setStyleSheet(
            """
            background: #F8F9FA;
            border-radius: 24px;
            border: 2px solid #E0E6ED;
            """
        )
        host_layout = QVBoxLayout(host_frame)
        host_title = QLabel("Принимающая сторона")
        host_title.setProperty("role", "block-title")
        host_title.setStyleSheet(
            """
            background: #FF9800;
            color: #fff;
            font-size: 22px;
            font-weight: bold;
            border-radius: 16px;
            padding: 12px 36px;
            margin-bottom: 24px;
            """
        )
        host_layout.addWidget(host_title)

        host_grid = QGridLayout()
        host_grid.setSpacing(20)

        self.department = QLineEdit()
        self.department.setPlaceholderText("Подразделение")
        self.department.setStyleSheet(
            """
            QLineEdit {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            """
        )
        self.host_fio = QLineEdit()
        self.host_fio.setPlaceholderText("ФИО принимающей стороны")
        self.host_fio.setStyleSheet(
            """
            QLineEdit {
                padding: 12px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 16px;
                background: white;
            }
            """
        )

        host_grid.addWidget(QLabel("Подразделение:"), 0, 0)
        host_grid.addWidget(self.department, 0, 1)
        host_grid.addWidget(QLabel("ФИО принимающей стороны:"), 1, 0)
        host_grid.addWidget(self.host_fio, 1, 1)
        host_layout.addLayout(host_grid)
        form_layout.addWidget(host_frame)
        visitors_frame = QFrame()
        visitors_frame.setStyleSheet(
            """
            background: #F8F9FA;
            border-radius: 24px;
            border: 2px solid #E0E6ED;
            """
        )
        visitors_layout = QVBoxLayout(visitors_frame)
        visitors_title = QLabel("Список посетителей")
        visitors_title.setProperty("role", "block-title")
        visitors_title.setStyleSheet(
            """
            background: #FF9800;
            color: #fff;
            font-size: 22px;
            font-weight: bold;
            border-radius: 16px;
            padding: 12px 36px;
            margin-bottom: 24px;
            """
        )
        visitors_layout.addWidget(visitors_title)
        self.visitors_table = QTableWidget()
        self.visitors_table.setColumnCount(10)
        self.visitors_table.setHorizontalHeaderLabels(
            [
                "Фамилия",
                "Имя",
                "Отчество",
                "Телефон",
                "Email",
                "Организация",
                "Примечание",
                "Дата рождения",
                "Серия",
                "Номер",
            ]
        )
        self.visitors_table.setStyleSheet(
            """
            QTableWidget {
                background: white;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                gridline-color: #E0E6ED;
            }
            QHeaderView::section {
                background: #F8F9FA;
                padding: 12px;
                border: none;
                border-bottom: 2px solid #E0E6ED;
                font-size: 16px;
                font-weight: bold;
                color: #333;
            }
            QTableWidget::item {
                padding: 12px;
                border-bottom: 1px solid #E0E6ED;
            }
            """
        )
        self.visitors_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        visitors_layout.addWidget(self.visitors_table)
        visitors_btns = QHBoxLayout()

        add_visitor_btn = QPushButton("Добавить посетителя")
        add_visitor_btn.setStyleSheet(
            """
            QPushButton {
                padding: 18px 0;
                border: 2px solid #FF9800;
                border-radius: 12px;
                font-size: 18px;
                font-weight: bold;
                color: #FF9800;
                background: white;
                min-width: 0;
            }
            QPushButton:hover {
                background: #FF9800;
                color: white;
            }
            """
        )
        add_visitor_btn.setCursor(Qt.PointingHandCursor)
        add_visitor_btn.clicked.connect(self.add_visitor)
        add_visitor_btn.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)

        import_visitors_btn = QPushButton("Импорт из Excel")
        import_visitors_btn.setStyleSheet(
            """
            QPushButton {
                padding: 18px 0;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 18px;
                font-weight: bold;
                color: #666;
                background: white;
                min-width: 0;
            }
            QPushButton:hover {
                background: #F8F9FA;
            }
            """
        )
        import_visitors_btn.setCursor(Qt.PointingHandCursor)
        import_visitors_btn.clicked.connect(self.import_visitors)
        import_visitors_btn.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)

        download_template_btn = QPushButton("Скачать шаблон")
        download_template_btn.setStyleSheet(
            """
            QPushButton {
                padding: 18px 0;
                border: 2px solid #7BA89F;
                border-radius: 12px;
                font-size: 18px;
                font-weight: bold;
                color: #fff;
                background: #7BA89F;
                min-width: 0;
            }
            QPushButton:hover {
                background: #5e8c7a;
                color: #fff;
            }
            """
        )
        download_template_btn.setCursor(Qt.PointingHandCursor)
        download_template_btn.clicked.connect(self.download_guest_template)
        download_template_btn.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)

        visitors_btns.addWidget(add_visitor_btn)
        visitors_btns.addWidget(import_visitors_btn)
        visitors_btns.addWidget(download_template_btn)
        visitors_layout.addLayout(visitors_btns)

        self.visitors_table.setMinimumHeight(400)
        self.visitors_table.setMaximumHeight(700)

        visitors_layout.addWidget(self.visitors_table)
        form_layout.addWidget(visitors_frame)

        docs_frame = QFrame()
        docs_frame.setStyleSheet(
            """
            background: #F8F9FA;
            border-radius: 24px;
            border: 2px solid #E0E6ED;
            """
        )
        docs_layout = QVBoxLayout(docs_frame)
        docs_title = QLabel("Документы")
        docs_title.setProperty("role", "block-title")
        docs_title.setStyleSheet(
            """
            background: #FF9800;
            color: #fff;
            font-size: 22px;
            font-weight: bold;
            border-radius: 16px;
            padding: 12px 36px;
            margin-bottom: 24px;
            """
        )
        docs_layout.addWidget(docs_title)

        docs_grid = QGridLayout()
        docs_grid.setSpacing(20)

        self.attach_btn = QPushButton("Прикрепить файл")
        self.attach_btn.setStyleSheet(
            """
            QPushButton {
                padding: 12px 24px;
                border: 2px solid #FF9800;
                border-radius: 12px;
                font-size: 16px;
                font-weight: bold;
                color: #FF9800;
                background: white;
            }
            QPushButton:hover {
                background: #FF9800;
                color: white;
            }
            """
        )
        self.attach_btn.setCursor(Qt.PointingHandCursor)
        self.attach_btn.clicked.connect(self.attach_file)

        docs_grid.addWidget(QLabel("Дополнительный файл:"), 0, 0)
        docs_grid.addWidget(self.attach_btn, 0, 1)

        docs_layout.addLayout(docs_grid)
        form_layout.addWidget(docs_frame)

        card_layout.addLayout(form_layout)

        btns = QHBoxLayout()
        clear_btn = QPushButton("Очистить форму")
        clear_btn.setStyleSheet(
            """
            QPushButton {
                padding: 16px 32px;
                border: 2px solid #E0E6ED;
                border-radius: 12px;
                font-size: 18px;
                font-weight: bold;
                color: #666;
                background: white;
            }
            QPushButton:hover {
                background: #F8F9FA;
            }
            """
        )
        clear_btn.setCursor(Qt.PointingHandCursor)
        clear_btn.clicked.connect(self.clear_group_form)

        submit_btn = QPushButton("Оформить заявку")
        submit_btn.setStyleSheet(
            """
            QPushButton {
                padding: 16px 32px;
                border: none;
                border-radius: 12px;
                font-size: 18px;
                font-weight: bold;
                color: white;
                background: #FF9800;
            }
            QPushButton:hover {
                background: #F57C00;
            }
            """
        )
        submit_btn.setCursor(Qt.PointingHandCursor)
        submit_btn.clicked.connect(self.submit_group_form)

        btns.addWidget(clear_btn)
        btns.addWidget(submit_btn)
        card_layout.addLayout(btns)

        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)
        scroll_area.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        scroll_area.setWidget(card)
        scroll_area.setStyleSheet("border: none; background: transparent;")
        main_layout.addWidget(scroll_area)

    def _normalize_header(self, header):
        header = (
            header.lower()
            .replace("ё", "е")
            .replace(" ", "")
            .replace("-", "")
            .replace("(", "")
            .replace(")", "")
            .replace(".", "")
        )
        if "фамил" in header:
            return "last_name"
        if "имя" in header:
            return "first_name"
        if "отчест" in header:
            return "middle_name"
        if "телефон" in header:
            return "phone"
        if "email" in header:
            return "email"
        if "организац" in header:
            return "org"
        if "примечан" in header:
            return "note"
        if "датарожд" in header:
            return "birth"
        if "серия" in header:
            return "series"
        if "номер" in header:
            return "number"
        return None

    def open_add_guest_dialog(self):
        from PyQt5.QtWidgets import QGridLayout
        from PyQt5.QtCore import QDate

        dialog = QDialog(self)
        dialog.setWindowTitle("Добавить гостя")
        dialog.setMinimumWidth(600)
        layout = QVBoxLayout(dialog)
        header = QLabel("Информация о посетителе")
        header.setAlignment(Qt.AlignCenter)
        header.setStyleSheet(
            "background:#FF9800;color:white;font-size:20px;font-weight:bold;border-top-left-radius:16px;border-top-right-radius:16px;padding:12px 0;margin-bottom:8px;"
        )
        layout.addWidget(header)
        form_grid = QGridLayout()
        form_grid.setHorizontalSpacing(18)
        form_grid.setVerticalSpacing(10)
        fields = {}
        form_grid.addWidget(QLabel("Фамилия*:"), 0, 0)
        fields["Фамилия"] = QLineEdit()
        form_grid.addWidget(fields["Фамилия"], 0, 1)
        form_grid.addWidget(QLabel("Имя*:"), 1, 0)
        fields["Имя"] = QLineEdit()
        form_grid.addWidget(fields["Имя"], 1, 1)
        form_grid.addWidget(QLabel("Отчество:"), 2, 0)
        fields["Отчество"] = QLineEdit()
        form_grid.addWidget(fields["Отчество"], 2, 1)
        form_grid.addWidget(QLabel("Телефон:"), 3, 0)
        fields["Телефон"] = QLineEdit()
        fields["Телефон"].setInputMask("+7 (000) 000-00-00;_")
        form_grid.addWidget(fields["Телефон"], 3, 1)
        form_grid.addWidget(QLabel("E-mail*:"), 4, 0)
        fields["Email"] = QLineEdit()
        form_grid.addWidget(fields["Email"], 4, 1)
        form_grid.addWidget(QLabel("Организация:"), 0, 2)
        fields["Организация"] = QLineEdit()
        form_grid.addWidget(fields["Организация"], 0, 3)
        form_grid.addWidget(QLabel("Примечание*:"), 1, 2)
        fields["Примечание"] = QLineEdit()
        form_grid.addWidget(fields["Примечание"], 1, 3)
        form_grid.addWidget(QLabel("Дата рождения*:"), 2, 2)
        fields["Дата рождения"] = QDateEdit(QDate.currentDate())
        fields["Дата рождения"].setCalendarPopup(True)
        fields["Дата рождения"].setDisplayFormat("dd.MM.yyyy")
        form_grid.addWidget(fields["Дата рождения"], 2, 3)
        form_grid.addWidget(QLabel("Серия*:"), 3, 2)
        fields["Серия"] = QLineEdit()
        fields["Серия"].setMaxLength(4)
        form_grid.addWidget(fields["Серия"], 3, 3)
        form_grid.addWidget(QLabel("Номер*:"), 4, 2)
        fields["Номер"] = QLineEdit()
        fields["Номер"].setMaxLength(6)
        form_grid.addWidget(fields["Номер"], 4, 3)
        layout.addLayout(form_grid)
        btns = QHBoxLayout()
        btns.addStretch()
        add_btn = QPushButton("Добавить посетителя")
        add_btn.setStyleSheet(
            "font-size:16px;padding:8px 32px;border-radius:8px;border:2px solid #FF9800;color:#FF9800;background:white;font-weight:bold;"
        )
        btns.addWidget(add_btn)
        layout.addLayout(btns)
        add_btn.clicked.connect(lambda: self.add_group_visitor_row_custom(fields))
        dialog.exec_()

    def download_guest_template(self):
        try:
            from openpyxl import Workbook
            from PyQt5.QtWidgets import QFileDialog

            wb = Workbook()
            ws = wb.active
            ws.title = "Гости"
            ws.append(
                [
                    "Фамилия",
                    "Имя",
                    "Отчество",
                    "Телефон (10 цифр, без +7)",
                    "Email",
                    "Организация",
                    "Примечание",
                    "Дата рождения (дд.мм.гггг)",
                    "Серия (4 цифры)",
                    "Номер (6 цифр)",
                ]
            )
            from openpyxl.comments import Comment

            ws["D1"].comment = Comment("Только 10 цифр, например: 9121234567", "AI")
            ws["I1"].comment = Comment("Формат: дд.мм.гггг", "AI")
            ws["J1"].comment = Comment("4 цифры", "AI")
            ws["K1"].comment = Comment("6 цифр", "AI")
            file_path, _ = QFileDialog.getSaveFileName(
                self, "Сохранить шаблон", "шаблон_гостей.xlsx", "Excel Files (*.xlsx)"
            )
            if file_path:
                wb.save(file_path)
                QMessageBox.information(self, "Успех", "Шаблон успешно сохранён!")
        except ImportError:
            QMessageBox.warning(
                self, "Ошибка", "Установите библиотеку openpyxl: pip install openpyxl"
            )

    def upload_guest_list(self):
        try:
            from openpyxl import load_workbook
            from PyQt5.QtWidgets import QFileDialog
            from PyQt5.QtCore import QDate
            import re

            file_path, _ = QFileDialog.getOpenFileName(
                self, "Загрузить список гостей", "", "Excel Files (*.xlsx)"
            )
            if not file_path:
                return
            wb = load_workbook(file_path)
            ws = wb.active
            headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]
            norm_map = {i: self._normalize_header(h) for i, h in enumerate(headers)}
            required = [
                "last_name",
                "first_name",
                "middle_name",
                "phone",
                "email",
                "org",
                "note",
                "birth",
                "series",
                "number",
            ]
            if not all(v in norm_map.values() for v in required):
                missing = [v for v in required if v not in norm_map.values()]
                QMessageBox.warning(
                    self, "Ошибка", f"В файле отсутствуют столбцы: {', '.join(missing)}"
                )
                return
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            for row in rows:
                if not row:
                    continue
                fields = {}
                for idx, norm in norm_map.items():
                    if norm:
                        fields[norm] = (
                            str(row[idx]).strip() if row[idx] is not None else ""
                        )
                phone_digits = "".join(filter(str.isdigit, fields.get("phone", "")))
                phone_formatted = ""
                if len(phone_digits) == 10:
                    phone_formatted = f"+7 ({phone_digits[0:3]}) {phone_digits[3:6]}-{phone_digits[6:8]}-{phone_digits[8:10]}"
                fields["phone"] = phone_formatted
                try:
                    d, m, y = map(int, str(fields.get("birth", "")).split("."))
                    birth_date = QDate(y, m, d)
                    fields["birth"] = birth_date.toString("dd.MM.yyyy")
                except Exception:
                    fields["birth"] = ""
                self.add_group_visitor_row_custom(fields)
            QMessageBox.information(self, "Успех", "Гости успешно загружены из файла!")
        except ImportError:
            QMessageBox.warning(
                self, "Ошибка", "Установите библиотеку openpyxl: pip install openpyxl"
            )
        except Exception as e:
            QMessageBox.warning(self, "Ошибка", f"Ошибка загрузки файла: {e}")

    def attach_file(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Выберите файл",
            "",
            "Все файлы (*.*);;Документы (*.pdf *.doc *.docx);;Изображения (*.jpg *.jpeg *.png)",
        )

        if file_path:
            self.attach_btn.setText(os.path.basename(file_path))
            self.attach_btn.setProperty("file_path", file_path)

    def clear_group_form(self):
        self.group_pass_id.setText(str(random.randint(100000, 999999)))
        self.group_date_from.setDate(QDate.currentDate())
        self.group_date_to.setDate(QDate.currentDate())
        self.group_purpose.setCurrentIndex(0)
        self.group_purpose_other.clear()
        self.group_purpose_other.hide()
        self.department.clear()
        self.host_fio.clear()
        self.visitors_table.setRowCount(0)

    def submit_group_form(self):
        errors = []
        if not self.department.text().strip():
            errors.append("Подразделение обязательно")
        if not self.host_fio.text().strip():
            errors.append("ФИО принимающей стороны обязательно")
        if self.visitors_table.rowCount() == 0:
            errors.append("Добавьте хотя бы одного посетителя в группу")

        if errors:
            QMessageBox.warning(self, "Ошибка", "\n".join(errors))
            return

        try:
            request_id = str(uuid.uuid4())
            pass_id = str(random.randint(100000, 999999))
            created_at = QDateTime.currentDateTime().toString(Qt.ISODate)
            visitors = []
            for row in range(self.visitors_table.rowCount()):
                visitor = {
                    "fio": (
                        self.visitors_table.item(row, 0).text()
                        if self.visitors_table.item(row, 0)
                        else ""
                    ),
                    "contacts": (
                        self.visitors_table.item(row, 1).text()
                        if self.visitors_table.item(row, 1)
                        else ""
                    ),
                    "birth": (
                        self.visitors_table.item(row, 2).text()
                        if self.visitors_table.item(row, 2)
                        else ""
                    ),
                    "passport": (
                        self.visitors_table.item(row, 3).text()
                        if self.visitors_table.item(row, 3)
                        else ""
                    ),
                    "photo_path": (
                        self.visitors_table.item(row, 4).text()
                        if self.visitors_table.item(row, 4)
                        else ""
                    ),
                }
                visitors.append(visitor)
            req = {
                "id": request_id,
                "type": "group",
                "status": "pending",
                "created_at": created_at,
                "pass_id": pass_id,
                "date_from": self.group_date_from.date().toString("yyyy-MM-dd"),
                "date_to": self.group_date_to.date().toString("yyyy-MM-dd"),
                "purpose": (
                    self.group_purpose_other.text()
                    if self.group_purpose.currentText() == "Другое"
                    else self.group_purpose.currentText()
                ),
                "department": self.department.text(),
                "host_fio": self.host_fio.text(),
                "visitors": visitors,
                "attached_file_path": (
                    self.attach_btn.property("file_path")
                    if self.attach_btn.property("file_path")
                    else ""
                ),
            }
            data = self.parent().load_requests()
            if not isinstance(data, list):
                data = []
            data.append(req)
            self.parent().save_requests(data)
            if hasattr(self.parent(), "show_requests"):
                self.parent().show_requests()

            QMessageBox.information(
                self,
                "Успех",
                f"Групповая заявка №{pass_id} успешно отправлена администратору на одобрение!",
            )
            self.clear_group_form()
            self.accept()

        except Exception as e:
            QMessageBox.critical(
                self, "Ошибка", f"Произошла ошибка при сохранении заявки: {str(e)}"
            )
            print(f"Ошибка при сохранении групповой заявки: {e}")

    def add_group_visitor_row_custom(self, fields):
        row_count = self.visitors_table.rowCount()
        if row_count < 8:
            self.visitors_table.insertRow(row_count)
            self.visitors_table.setItem(
                row_count, 0, QTableWidgetItem(fields.get("last_name", ""))
            )
            self.visitors_table.setItem(
                row_count, 1, QTableWidgetItem(fields.get("first_name", ""))
            )
            self.visitors_table.setItem(
                row_count, 2, QTableWidgetItem(fields.get("middle_name", ""))
            )
            self.visitors_table.setItem(
                row_count, 3, QTableWidgetItem(fields.get("phone", ""))
            )
            self.visitors_table.setItem(
                row_count, 4, QTableWidgetItem(fields.get("email", ""))
            )
            self.visitors_table.setItem(
                row_count, 5, QTableWidgetItem(fields.get("org", ""))
            )
            self.visitors_table.setItem(
                row_count, 6, QTableWidgetItem(fields.get("note", ""))
            )
            self.visitors_table.setItem(
                row_count, 7, QTableWidgetItem(fields.get("birth", ""))
            )
            self.visitors_table.setItem(
                row_count, 8, QTableWidgetItem(fields.get("series", ""))
            )
            self.visitors_table.setItem(
                row_count, 9, QTableWidgetItem(fields.get("number", ""))
            )
            delete_btn = QPushButton("Удалить")
            delete_btn.clicked.connect(lambda: self.delete_visitor_row(row_count))
            self.visitors_table.setCellWidget(row_count, 10, delete_btn)
        else:
            QMessageBox.warning(
                self, "Внимание", "Максимальное количество посетителей в группе - 8."
            )

    def delete_visitor_row(self, row):
        self.visitors_table.removeRow(row)

    def on_purpose_changed(self, text):
        self.group_purpose_other.setVisible(text == "Другое")
        if text == "Другое":
            self.group_purpose_other.setFocus()

    def add_visitor(self):
        from PyQt5.QtWidgets import (
            QDialog,
            QVBoxLayout,
            QFormLayout,
            QLineEdit,
            QDateEdit,
            QPushButton,
            QHBoxLayout,
            QMessageBox,
        )
        from PyQt5.QtCore import QDate
        from PyQt5.QtGui import QRegExpValidator
        from PyQt5.QtCore import QRegExp

        dialog = QDialog(self)
        dialog.setWindowTitle("Добавить посетителя")
        layout = QVBoxLayout(dialog)
        form = QFormLayout()
        last_name = QLineEdit()
        first_name = QLineEdit()
        middle_name = QLineEdit()
        phone = QLineEdit()
        email = QLineEdit()
        org = QLineEdit()
        note = QLineEdit()
        birth = QDateEdit(QDate.currentDate())
        birth.setCalendarPopup(True)
        series = QLineEdit()
        number = QLineEdit()
        form.addRow("Фамилия*:", last_name)
        form.addRow("Имя*:", first_name)
        form.addRow("Отчество:", middle_name)
        form.addRow("Телефон:", phone)
        form.addRow("Email*:", email)
        form.addRow("Организация:", org)
        form.addRow("Примечание*:", note)
        form.addRow("Дата рождения*:", birth)
        form.addRow("Серия*:", series)
        form.addRow("Номер*:", number)
        layout.addLayout(form)
        btns = QHBoxLayout()
        save_btn = QPushButton("Сохранить")
        cancel_btn = QPushButton("Отмена")
        btns.addWidget(save_btn)
        btns.addWidget(cancel_btn)
        layout.addLayout(btns)

        def save():
            if (
                not last_name.text().strip()
                or not first_name.text().strip()
                or not email.text().strip()
                or not note.text().strip()
                or not series.text().strip()
                or not number.text().strip()
            ):
                QMessageBox.warning(
                    dialog,
                    "Ошибка",
                    "Пожалуйста, заполните все обязательные поля со звёздочкой (*)",
                )
                return
            fields = {
                "last_name": last_name.text(),
                "first_name": first_name.text(),
                "middle_name": middle_name.text(),
                "phone": phone.text(),
                "email": email.text(),
                "org": org.text(),
                "note": note.text(),
                "birth": birth.date().toString("dd.MM.yyyy"),
                "series": series.text(),
                "number": number.text(),
            }
            self.add_group_visitor_row_custom(fields)
            dialog.accept()

        save_btn.clicked.connect(save)
        cancel_btn.clicked.connect(dialog.reject)
        dialog.exec_()

    def import_visitors(self):
        try:
            from openpyxl import load_workbook
            from PyQt5.QtWidgets import QFileDialog
            from PyQt5.QtCore import QDate
            import re

            file_path, _ = QFileDialog.getOpenFileName(
                self,
                "Выберите файл со списком посетителей",
                "",
                "Excel Files (*.xlsx *.xls)",
            )
            if not file_path:
                return
            wb = load_workbook(file_path)
            ws = wb.active
            headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]
            norm_map = {i: self._normalize_header(h) for i, h in enumerate(headers)}
            required = [
                "last_name",
                "first_name",
                "middle_name",
                "phone",
                "email",
                "org",
                "note",
                "birth",
                "series",
                "number",
            ]
            if not all(v in norm_map.values() for v in required):
                missing = [v for v in required if v not in norm_map.values()]
                QMessageBox.warning(
                    self, "Ошибка", f"В файле отсутствуют столбцы: {', '.join(missing)}"
                )
                return
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            for row in rows:
                if not row:
                    continue
                fields = {}
                for idx, norm in norm_map.items():
                    if norm:
                        fields[norm] = (
                            str(row[idx]).strip() if row[idx] is not None else ""
                        )

                phone_digits = "".join(filter(str.isdigit, fields.get("phone", "")))
                phone_formatted = ""
                if len(phone_digits) == 10:
                    phone_formatted = f"+7 ({phone_digits[0:3]}) {phone_digits[3:6]}-{phone_digits[6:8]}-{phone_digits[8:10]}"
                fields["phone"] = phone_formatted
                try:
                    d, m, y = map(int, str(fields.get("birth", "")).split("."))
                    birth_date = QDate(y, m, d)
                    fields["birth"] = birth_date.toString("dd.MM.yyyy")
                except Exception:
                    fields["birth"] = ""
                self.add_group_visitor_row_custom(fields)
            QMessageBox.information(self, "Успех", "Гости успешно загружены из файла!")
        except ImportError:
            QMessageBox.warning(
                self, "Ошибка", "Установите библиотеку openpyxl: pip install openpyxl"
            )
        except Exception as e:
            QMessageBox.warning(self, "Ошибка", f"Ошибка загрузки файла: {e}")


class GuardPanel(QMainWindow):
    def __init__(self, fio):
        super().__init__()
        self.fio = fio
        self.setWindowTitle("Guard")
        self.setMinimumSize(1200, 800)
        self.init_ui()

    def init_ui(self):
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QHBoxLayout(central_widget)  
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        menu_widget = QWidget()
        menu_widget.setStyleSheet(
            """
            background: qlineargradient(x1:0, y1:0, x2:1, y2:1, stop:0 #e0eafc, stop:1 #cfdef3);
            border-top-left-radius: 32px;
            border-bottom-left-radius: 32px;
            box-shadow: 6px 0 32px #b0b0b033;
            """
        )
        menu_layout = QVBoxLayout(menu_widget)
        menu_layout.setAlignment(Qt.AlignTop)
        menu_widget.setFixedWidth(320)
        self.events_btn = QPushButton()
        self.requests_btn = QPushButton()
        self.create_request_btn = QPushButton()
        from PyQt5.QtGui import QPixmap

        avatar = QLabel()
        avatar.setPixmap(QPixmap(64, 64))
        avatar.setStyleSheet(
            "background:#fff;border-radius:32px;border:3px solid #7BA89F;margin-top:32px;margin-bottom:8px;"
        )
        avatar.setFixedSize(64, 64)
        menu_layout.addWidget(avatar, alignment=Qt.AlignHCenter)
        fio_lbl = QLabel(f"<b style='font-size:18px;color:#7BA89F;'>{self.fio}</b>")
        fio_lbl.setAlignment(Qt.AlignCenter)
        fio_lbl.setStyleSheet("margin-bottom:18px;")
        menu_layout.addWidget(fio_lbl)
        title = QLabel(
            "<span style='font-size:28px;font-weight:900;letter-spacing:2px;color:#FF9800;text-shadow:0 2px 8px #7BA89F;'>Вахтер</span>"
        )
        title.setAlignment(Qt.AlignCenter)
        title.setStyleSheet("margin-bottom: 18px; padding: 0 10px;")
        menu_layout.addWidget(title)
        sep = QFrame()
        sep.setFrameShape(QFrame.HLine)
        sep.setStyleSheet("color:#e0e6ed;margin:0 0 18px 0;")
        menu_layout.addWidget(sep)
        btn_style = (
            "QPushButton {"
            "background: #fff; font-size: 17px; padding: 14px 0; margin-bottom: 18px; border-radius: 14px; border: none;"
            "box-shadow: 0 2px 8px #b0b0b022; color: #222; font-weight: 600; letter-spacing: 1px;"
            "transition: background 0.3s, color 0.3s;"
            "text-align: left; padding-left: 32px;"
            "}"
            "QPushButton:hover {"
            "background: #FF9800; color: #fff;"
            "}"
        )
        menu_buttons = [
            (self.events_btn, "📋 Журнал событий"),
            (self.requests_btn, "📄 Заявки"),
            (self.create_request_btn, "➕ Создать заявку"),
        ]
        for btn, text in menu_buttons:
            btn.setText(text)
            btn.setStyleSheet(btn_style)
            btn.setCursor(Qt.PointingHandCursor)
            btn.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
            menu_layout.addWidget(btn)

        self.events_btn.clicked.connect(self.show_events)
        self.requests_btn.clicked.connect(self.show_requests)
        self.create_request_btn.clicked.connect(self.show_create_request_tiles)

        menu_layout.addStretch()
        sep2 = QFrame()
        sep2.setFrameShape(QFrame.HLine)
        sep2.setStyleSheet("color:#e0e6ed;margin:18px 0 0 0;")
        menu_layout.addWidget(sep2)
        version_lbl = QLabel("<span style='color:#bbb;font-size:13px;'>v1.0</span>")
        version_lbl.setAlignment(Qt.AlignCenter)
        version_lbl.setStyleSheet("margin-top:8px;margin-bottom:8px;")
        menu_layout.addWidget(version_lbl)
        logout_btn = QPushButton("Выйти")
        logout_btn.setStyleSheet(
            "background:#eee;font-size:15px;padding:8px 0;border-radius:10px;color:#888;font-weight:600;"
        )
        logout_btn.setCursor(Qt.PointingHandCursor)
        menu_layout.addWidget(logout_btn)
        main_layout.addWidget(menu_widget)
        self.work_area = QWidget()
        self.work_area.setStyleSheet(
            """
            background: #F3F4F6;
            border-top-right-radius: 24px;
            border-bottom-right-radius: 24px;
            """
        )
        self.work_layout = QVBoxLayout(self.work_area)
        main_layout.addWidget(self.work_area)

    def show_events(self):
        self.clear_work_area()
        title = QLabel("Журнал событий")
        title.setStyleSheet(
            """
            font-size: 32px;
            font-weight: bold;
            color: #2c3e50;
            margin-bottom: 20px;
        """
        )
        self.work_layout.addWidget(title)
        events_table = QTableWidget()
        events_table.setStyleSheet(
            """
            QTableWidget {
                background: white;
                border-radius: 16px;
                border: none;
                gridline-color: #E0E6ED;
            }
            QTableWidget::item {
                padding: 8px;
                border-bottom: 1px solid #E0E6ED;
            }
            QHeaderView::section {
                background: #F8F9FA;
                padding: 12px;
                border: none;
                border-bottom: 2px solid #E0E6ED;
                font-weight: bold;
                color: #2c3e50;
            }
        """
        )
        events_table.setColumnCount(3)
        events_table.setHorizontalHeaderLabels(["Время", "Событие", "Статус"])
        events_table.horizontalHeader().setSectionResizeMode(
            0, QHeaderView.ResizeToContents
        )
        events_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        events_table.horizontalHeader().setSectionResizeMode(
            2, QHeaderView.ResizeToContents
        )
        events_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        events_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        events_table.verticalHeader().setVisible(False)
        try:
            with open("traffic_history.json", "r", encoding="utf-8") as f:
                history = json.load(f)
                events_table.setRowCount(len(history))
                for i, event in enumerate(history):
                    time_item = QTableWidgetItem(event.get("time", ""))
                    message_item = QTableWidgetItem(event.get("message", ""))
                    status_item = QTableWidgetItem(
                        "Вход" if "Вход" in event.get("message", "") else "Выход"
                    )
                    status_item.setForeground(
                        QColor(
                            "#27ae60"
                            if "Вход" in event.get("message", "")
                            else "#e74c3c"
                        )
                    )

                    events_table.setItem(i, 0, time_item)
                    events_table.setItem(i, 1, message_item)
                    events_table.setItem(i, 2, status_item)
        except FileNotFoundError:
            pass

        self.work_layout.addWidget(events_table)

    def show_requests(self):
        self.clear_work_area()
        title = QLabel("Заявки на посещение")
        title.setStyleSheet(
            """
            QLabel {
                font-size: 32px;
                font-weight: bold;
                color: #2c3e50;
                margin-bottom: 20px;
            }
        """
        )
        self.work_layout.addWidget(title)
        self.request_tabs = QTabWidget()
        self.request_tabs.setStyleSheet(
            """
            QTabWidget::pane {
                border: 0;
                background: #F3F4F6;
                border-radius: 16px;
                margin-top: 12px;
            }
            QTabBar::tab {
                background: #fff;
                color: #000;
                padding: 12px 24px;
                border-top-left-radius: 12px;
                border-top-right-radius: 12px;
                margin-right: 4px;
                font-size: 16px;
                font-weight: bold;
                min-width: 160px;
            }
            QTabBar::tab:selected {
                background: #FF9800;
                color: #000;
            }
            QTabBar::tab:hover:!selected {
                background: #ffe0b2;
                color: #000;
            }
        """
        )
        all_requests_tab = QWidget()
        all_requests_layout = QVBoxLayout(all_requests_tab)
        self.all_requests_table = self.create_requests_table()
        all_requests_layout.addWidget(self.all_requests_table)
        self.request_tabs.addTab(all_requests_tab, "Все заявки")
        approved_requests_tab = QWidget()
        approved_requests_layout = QVBoxLayout(approved_requests_tab)
        self.approved_requests_table = self.create_requests_table()
        approved_requests_layout.addWidget(self.approved_requests_table)
        self.request_tabs.addTab(approved_requests_tab, "Одобренные")
        rejected_requests_tab = QWidget()
        rejected_requests_layout = QVBoxLayout(rejected_requests_tab)
        self.rejected_requests_table = self.create_requests_table()
        rejected_requests_layout.addWidget(self.rejected_requests_table)
        self.request_tabs.addTab(rejected_requests_tab, "Отклоненные")
        pending_requests_tab = QWidget()
        pending_requests_layout = QVBoxLayout(pending_requests_tab)
        self.pending_requests_table = self.create_requests_table()
        pending_requests_layout.addWidget(self.pending_requests_table)
        self.request_tabs.addTab(pending_requests_tab, "В процессе")
        history_tab = QWidget()
        history_layout = QVBoxLayout(history_tab)
        filter_layout = QHBoxLayout()
        date_filter_label = QLabel("Период:")
        date_filter_label.setStyleSheet("font-size: 14px; color: #666;")
        date_from = QDateEdit()
        date_from.setCalendarPopup(True)
        date_from.setDisplayFormat("dd.MM.yyyy")
        date_from.setDate(QDate.currentDate().addMonths(-1))
        date_to = QDateEdit()
        date_to.setCalendarPopup(True)
        date_to.setDisplayFormat("dd.MM.yyyy")
        date_to.setDate(QDate.currentDate())
        type_filter_label = QLabel("Тип заявки:")
        type_filter_label.setStyleSheet("font-size: 14px; color: #666;")
        type_filter = QComboBox()
        type_filter.addItems(["Все", "Индивидуальная", "Групповая"])
        apply_filter_btn = QPushButton("Применить")
        apply_filter_btn.setStyleSheet(
            """
            QPushButton {
                background-color: #FF9800;
                color: white;
                border: none;
                padding: 8px 16px;
                border-radius: 8px;
                font-weight: bold;
            }
            QPushButton:hover {
                background-color: #F57C00;
            }
        """
        )

        filter_layout.addWidget(date_filter_label)
        filter_layout.addWidget(date_from)
        filter_layout.addWidget(QLabel("—"))
        filter_layout.addWidget(date_to)
        filter_layout.addSpacing(20)
        filter_layout.addWidget(type_filter_label)
        filter_layout.addWidget(type_filter)
        filter_layout.addSpacing(20)
        filter_layout.addWidget(apply_filter_btn)
        filter_layout.addStretch()

        history_layout.addLayout(filter_layout)

        self.history_table = self.create_history_table()
        history_layout.addWidget(self.history_table)

        self.request_tabs.addTab(history_tab, "История")

        self.work_layout.addWidget(self.request_tabs)

        self.update_requests_tables()

        apply_filter_btn.clicked.connect(
            lambda: self.apply_history_filters(
                date_from.date(), date_to.date(), type_filter.currentText()
            )
        )

    def create_history_table(self):

        table = QTableWidget()
        table.setColumnCount(8)
        table.setHorizontalHeaderLabels(
            [
                "№",
                "Тип заявки",
                "ФИО посетителя",
                "Дата создания",
                "Действует с",
                "Действует по",
                "Статус",
                "Действия",
            ]
        )

        table.setMinimumHeight(600)

        table.verticalHeader().setDefaultSectionSize(50)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setSelectionMode(QAbstractItemView.SingleSelection)
        table.setStyleSheet(
            """
            QTableWidget {
                background-color: white;
                border: 1px solid #e0e0e0;
                border-radius: 12px;
                gridline-color: #f0f0f0;
            }
            QTableWidget::item {
                padding: 12px;
                border-bottom: 1px solid #f0f0f0;
            }
            QHeaderView::section {
                background-color: #f8f9fa;
                padding: 16px;
                border: none;
                border-bottom: 2px solid #e0e0e0;
                font-weight: bold;
                color: #2c3e50;
            }
            QTableWidget::item:selected {
                background-color: #fff3e0;
                color: #FF9800;
            }
            QPushButton {
                background-color: #fff;
                color: #FF9800;
                border: 2px solid #FF9800;
                padding: 8px 24px;
                border-radius: 8px;
                font-size: 16px;
                min-width: 120px;
                font-weight: bold;
            }
            QPushButton:hover {
                background-color: #FF9800;
                color: #fff;
            }
        """
        )
        table.setColumnWidth(7, 180)  
        return table

    def apply_history_filters(self, date_from, date_to, request_type):
        requests = self.load_requests()
        filtered_requests = []

        for request in requests:

            created_at = request.get("created_at", "")
            if not created_at:
                continue

            try:
                created_date = QDateTime.fromString(created_at, Qt.ISODate).date()
                if not created_date.isValid():
                    continue

                if date_from <= created_date <= date_to:

                    if (
                        request_type == "Все"
                        or (
                            request_type == "Индивидуальная"
                            and request.get("type") == "individual"
                        )
                        or (
                            request_type == "Групповая"
                            and request.get("type") == "group"
                        )
                    ):
                        filtered_requests.append(request)
            except Exception as e:
                print(f"Ошибка при обработке даты: {e}")
                continue

        self.populate_history_table(filtered_requests)

    def populate_history_table(self, requests_list):
        self.history_table.setRowCount(len(requests_list))
        from PyQt5.QtGui import QIcon

        for row, request in enumerate(requests_list):
            id_item = QTableWidgetItem(str(request.get("id", "")))
            self.history_table.setItem(row, 0, id_item)

            type_text = (
                "Индивидуальная" if request.get("type") == "individual" else "Групповая"
            )
            type_item = QTableWidgetItem(type_text)
            self.history_table.setItem(row, 1, type_item)

            fio = f"{request.get('last_name', '')} {request.get('first_name', '')} {request.get('middle_name', '')}".strip()
            fio_item = QTableWidgetItem(fio)
            self.history_table.setItem(row, 2, fio_item)

            created_item = QTableWidgetItem(request.get("created_at", ""))
            self.history_table.setItem(row, 3, created_item)

            date_from_item = QTableWidgetItem(request.get("date_from", ""))
            self.history_table.setItem(row, 4, date_from_item)

            date_to_item = QTableWidgetItem(request.get("date_to", ""))
            self.history_table.setItem(row, 5, date_to_item)

            status = request.get("status", "pending")
            status_text = {
                "approved": "Одобрена",
                "rejected": "Отклонена",
                "pending": "В процессе",
            }.get(status, "В процессе")

            status_item = QTableWidgetItem(status_text)
            status_item.setForeground(self.get_status_color_brush(status))
            self.history_table.setItem(row, 6, status_item)

            view_btn = QPushButton("Просмотреть")
            view_btn.setStyleSheet(
                "font-size:12px;padding:2px 6px;border-radius:6px;background:#7BA89F;color:#000;"
            )
            view_btn.setProperty("action-btn", True)
            view_btn.setCursor(Qt.PointingHandCursor)
            view_btn.setIcon(QIcon.fromTheme("search") or QIcon("icons/search.png"))
            view_btn.setIconSize(QSize(20, 20))
            view_btn.clicked.connect(
                lambda checked, r=request: self.view_request_details(r)
            )

            cell_widget = QWidget()
            layout = QHBoxLayout(cell_widget)
            layout.addWidget(view_btn)
            layout.setAlignment(Qt.AlignCenter)
            layout.setContentsMargins(0, 0, 0, 0)
            cell_widget.setLayout(layout)
            cell_widget.setStyleSheet("background: #FFF3E0; border-radius: 12px;")
            self.history_table.setCellWidget(row, 7, cell_widget)

    def clear_work_area(self):
        while self.work_layout.count():
            item = self.work_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()

    def show_create_request_tiles(self):
        self.clear_work_area()
        main_container = QWidget()
        main_layout = QVBoxLayout(main_container)
        main_layout.setAlignment(Qt.AlignCenter)
        main_layout.setContentsMargins(50, 50, 50, 50)
        main_layout.setSpacing(40)
        title = QLabel("Создание заявки")
        title.setStyleSheet(
            """
            QLabel {
                font-size: 36px;
                font-weight: bold;
                color: #2c3e50;
                margin-bottom: 40px;
                text-align: center;
            }
        """
        )
        title.setAlignment(Qt.AlignCenter)
        main_layout.addWidget(title)
        tiles_container = QWidget()
        tiles_layout = QHBoxLayout(tiles_container)
        tiles_layout.setSpacing(40)
        tiles_layout.setContentsMargins(0, 0, 0, 0)
        tiles_layout.setAlignment(Qt.AlignCenter)
        individual_tile = QFrame()
        individual_tile.setObjectName("individualTile")
        individual_tile.setStyleSheet(
            """
            QFrame#individualTile {
                background-color: white;
                border-radius: 24px;
                border: 2px solid #e0e0e0;
                min-width: 450px;
                min-height: 550px;
                transition: all 0.3s ease;
            }
            QFrame#individualTile:hover {
                border-color: #FF9800;
                box-shadow: 0 8px 16px rgba(255, 152, 0, 0.15);
                transform: translateY(-5px);
            }
            QLabel {
                color: #2c3e50;
            }
            QPushButton {
                background-color: #FF9800;
                color: white;
                border: none;
                padding: 16px 32px;
                border-radius: 12px;
                font-size: 18px;
                font-weight: bold;
                min-width: 240px;
                transition: all 0.3s ease;
            }
            QPushButton:hover {
                background-color: #F57C00;
                transform: scale(1.05);
            }
        """
        )
        individual_layout = QVBoxLayout(individual_tile)
        individual_layout.setContentsMargins(40, 40, 40, 40)
        individual_layout.setSpacing(30)
        individual_layout.setAlignment(Qt.AlignCenter)
        individual_icon = QLabel()
        individual_icon.setPixmap(
            QPixmap("icons/individual.png").scaled(
                140, 140, Qt.KeepAspectRatio, Qt.SmoothTransformation
            )
        )
        individual_icon.setAlignment(Qt.AlignCenter)
        individual_layout.addWidget(individual_icon)
        individual_title = QLabel("Индивидуальная заявка")
        individual_title.setStyleSheet(
            "font-size: 28px; font-weight: bold; color: #2c3e50;"
        )
        individual_title.setAlignment(Qt.AlignCenter)
        individual_layout.addWidget(individual_title)
        individual_desc = QLabel(
            "Создание заявки для одного посетителя.\n"
            "Включает персональные данные и документы."
        )
        individual_desc.setStyleSheet("font-size: 18px; color: #666; line-height: 1.5;")
        individual_desc.setAlignment(Qt.AlignCenter)
        individual_desc.setWordWrap(True)
        individual_layout.addWidget(individual_desc)

        individual_layout.addStretch()
        individual_btn = QPushButton("Создать заявку")
        individual_btn.setCursor(Qt.PointingHandCursor)
        individual_btn.clicked.connect(self.open_individual_request_window)
        individual_layout.addWidget(individual_btn, alignment=Qt.AlignCenter)
        group_tile = QFrame()
        group_tile.setObjectName("groupTile")
        group_tile.setStyleSheet(
            """
            QFrame#groupTile {
                background-color: white;
                border-radius: 24px;
                border: 2px solid #e0e0e0;
                min-width: 450px;
                min-height: 550px;
                transition: all 0.3s ease;
            }
            QFrame#groupTile:hover {
                border-color: #FF9800;
                box-shadow: 0 8px 16px rgba(255, 152, 0, 0.15);
                transform: translateY(-5px);
            }
            QLabel {
                color: #2c3e50;
            }
            QPushButton {
                background-color: #FF9800;
                color: white;
                border: none;
                padding: 16px 32px;
                border-radius: 12px;
                font-size: 18px;
                font-weight: bold;
                min-width: 240px;
                transition: all 0.3s ease;
            }
            QPushButton:hover {
                background-color: #F57C00;
                transform: scale(1.05);
            }
        """
        )
        group_layout = QVBoxLayout(group_tile)
        group_layout.setContentsMargins(40, 40, 40, 40)
        group_layout.setSpacing(30)
        group_layout.setAlignment(Qt.AlignCenter)
        group_icon = QLabel()
        group_icon.setPixmap(
            QPixmap("icons/group.png").scaled(
                140, 140, Qt.KeepAspectRatio, Qt.SmoothTransformation
            )
        )
        group_icon.setAlignment(Qt.AlignCenter)
        group_layout.addWidget(group_icon)
        group_title = QLabel("Групповая заявка")
        group_title.setStyleSheet("font-size: 28px; font-weight: bold; color: #2c3e50;")
        group_title.setAlignment(Qt.AlignCenter)
        group_layout.addWidget(group_title)
        group_desc = QLabel(
            "Создание заявки для группы посетителей.\n"
            "Возможность загрузки списка из файла."
        )
        group_desc.setStyleSheet("font-size: 18px; color: #666; line-height: 1.5;")
        group_desc.setAlignment(Qt.AlignCenter)
        group_desc.setWordWrap(True)
        group_layout.addWidget(group_desc)

        group_layout.addStretch()
        group_btn = QPushButton("Создать заявку")
        group_btn.setCursor(Qt.PointingHandCursor)
        group_btn.clicked.connect(self.open_group_request_window)
        group_layout.addWidget(group_btn, alignment=Qt.AlignCenter)
        tiles_layout.addWidget(individual_tile)
        tiles_layout.addWidget(group_tile)
        main_layout.addWidget(tiles_container)
        main_layout.addStretch()
        self.work_layout.addWidget(main_container)
        self.work_layout.addStretch()

    def open_individual_request_window(self):
        dialog = CreateIndividualRequestWindow(self)
        dialog.exec_()

    def open_group_request_window(self):
        dialog = CreateGroupRequestWindow(self)
        dialog.exec_()

    def create_requests_table(self):

        table = QTableWidget()
        table.setColumnCount(7)
        table.setHorizontalHeaderLabels(
            [
                "№",
                "Тип заявки",
                "ФИО посетителя",
                "Дата создания",
                "Действует с",
                "Действует по",
                "Статус",
            ]
        )
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setSelectionMode(QAbstractItemView.SingleSelection)
        table.setStyleSheet(
            """
            QTableWidget {
                background-color: white;
                border: 1px solid #e0e0e0;
                border-radius: 12px;
                gridline-color: #f0f0f0;
            }
            QTableWidget::item {
                padding: 12px;
                border-bottom: 1px solid #f0f0f0;
            }
            QHeaderView::section {
                background-color: #f8f9fa;
                padding: 16px;
                border: none;
                border-bottom: 2px solid #e0e0e0;
                font-weight: bold;
                color: #2c3e50;
            }
            QTableWidget::item:selected {
                background-color: #fff3e0;
                color: #FF9800;
            }
        """
        )
        return table

    def update_requests_tables(self):
        requests = self.load_requests()
        approved_requests = [r for r in requests if r.get("status") == "approved"]
        rejected_requests = [r for r in requests if r.get("status") == "rejected"]
        pending_requests = [r for r in requests if r.get("status") == "pending"]
        self.populate_requests_table(self.all_requests_table, requests)
        self.populate_requests_table(self.approved_requests_table, approved_requests)
        self.populate_requests_table(self.rejected_requests_table, rejected_requests)
        self.populate_requests_table(self.pending_requests_table, pending_requests)

    def populate_requests_table(self, table, requests_list):
        table.setRowCount(len(requests_list))

        for row, request in enumerate(requests_list):

            id_item = QTableWidgetItem(str(request.get("id", "")))
            table.setItem(row, 0, id_item)
            type_text = (
                "Индивидуальная" if request.get("type") == "individual" else "Групповая"
            )
            type_item = QTableWidgetItem(type_text)
            table.setItem(row, 1, type_item)

            fio = f"{request.get('last_name', '')} {request.get('first_name', '')} {request.get('middle_name', '')}".strip()
            fio_item = QTableWidgetItem(fio)
            table.setItem(row, 2, fio_item)

            created_item = QTableWidgetItem(request.get("created_at", ""))
            table.setItem(row, 3, created_item)

            date_from_item = QTableWidgetItem(request.get("date_from", ""))
            table.setItem(row, 4, date_from_item)

            date_to_item = QTableWidgetItem(request.get("date_to", ""))
            table.setItem(row, 5, date_to_item)

            status = request.get("status", "pending")
            status_text = {
                "approved": "Одобрена",
                "rejected": "Отклонена",
                "pending": "В процессе",
            }.get(status, "В процессе")

            status_item = QTableWidgetItem(status_text)
            status_item.setForeground(self.get_status_color_brush(status))
            table.setItem(row, 6, status_item)

    def get_status_color_brush(self, status):
        from PyQt5.QtGui import QColor

        colors = {
            "approved": QColor("#4CAF50"),  
            "rejected": QColor("#F44336"),  
            "pending": QColor("#FF9800"),  
        }
        return colors.get(status, QColor("#FF9800"))

    def load_requests(self):
        try:
            if os.path.exists(REQUESTS_FILE):
                with open(REQUESTS_FILE, "r", encoding="utf-8") as f:
                    return json.load(f)
            return []
        except Exception as e:
            print(f"Ошибка при загрузке заявок: {e}")
            return []

    def save_requests(self, data):
        try:
            with open(REQUESTS_FILE, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print(f"Ошибка при сохранении заявок: {e}")

    def view_request_details(self, request):
        if request.get("type") == "individual":
            dialog = ViewEditIndividualRequestDialog(
                request, is_editable=False, parent=self
            )
        else:
            dialog = ViewEditGroupRequestDialog(request, is_editable=False, parent=self)
        dialog.exec_()
